from importlib.resources import files
import os
import json
import time
import uuid
import shutil
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, HTTPException, Query, Depends, Request, BackgroundTasks
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from app.database.connection import get_db, SessionLocal
from app.auth.jwt_handler import get_admin_user, get_current_active_user
from app.models.user_and_log import User, ReportJob
from app.services.audit_logger import log_audit
from utils.logger import logger
from config import settings
from models import UploadedFiles, ReportResponse
from services.validators import (
    validate_uploaded_files,
    validate_mapping_workbook,
    validate_budget_workbook,
    validate_trial_balance,
)
from services.file_reader import (
    read_current_year_tb,
    read_previous_year_tb,
    detect_tb_month_and_year,
)
from services.mapping_engine import MappingEngine
from services.tb_processor import process_trial_balance
from services.budget_processor import BudgetProcessor
from services.calculations import calculate_revenue
from services.ppt_generator import generate_pptx

router = APIRouter(prefix="/api", tags=["finance"])

uploaded_files_store: dict = {}
report_jobs: dict = {}


def _db_save_job(session_id: str, status: str, result: dict | None = None):
    """Upsert a report job row in the database (visible to all workers)."""
    db = SessionLocal()
    try:
        row = db.query(ReportJob).filter(ReportJob.session_id == session_id).first()
        if row:
            row.status = status
            row.result_json = json.dumps(result) if result is not None else None
        else:
            row = ReportJob(
                session_id=session_id,
                status=status,
                result_json=json.dumps(result) if result is not None else None,
            )
            db.add(row)
        db.commit()
    except Exception as exc:
        logger.warning(f"Could not persist job to DB for {session_id}: {exc}")
    finally:
        db.close()


def _db_load_job(session_id: str) -> dict | None:
    """Load a report job row from the database. Returns None if not found."""
    db = SessionLocal()
    try:
        row = db.query(ReportJob).filter(ReportJob.session_id == session_id).first()
        if row and row.result_json:
            return json.loads(row.result_json)
        if row and row.status == "processing":
            return {"status": "processing", "message": "Report generation in progress."}
        return None
    except Exception as exc:
        logger.warning(f"Could not load job from DB for {session_id}: {exc}")
        return None
    finally:
        db.close()


DEFAULT_FILES_DIR = settings.TEMPLATE_DIR / "defaults"
DEFAULT_FILES_DIR.mkdir(parents=True, exist_ok=True)


def _default_file_path(file_type: str) -> Path:
    if file_type not in {"budget", "mapping"}:
        raise HTTPException(status_code=400, detail="file_type must be 'budget' or 'mapping'.")
    matches = sorted(DEFAULT_FILES_DIR.glob(f"default_{file_type}.*"))
    return matches[-1] if matches else DEFAULT_FILES_DIR / f"default_{file_type}.xlsx"


def _default_file_info(file_type: str) -> dict:
    path = _default_file_path(file_type)
    return {
        f"default_{file_type}_active": path.exists(),
        f"default_{file_type}_filename": path.name if path.exists() else None,
    }


def _restore_session_files(session_id: str) -> dict | None:
    session_dir = settings.UPLOAD_DIR / session_id
    if not session_dir.exists() or not session_dir.is_dir():
        return None

    file_paths = {}
    for label in ["tb_current", "tb_previous", "budget", "mapping"]:
        matches = sorted(session_dir.glob(f"{label}__*"))
        if matches:
            file_paths[label] = str(matches[-1])

    for label in ["budget", "mapping"]:
        if label not in file_paths:
            default_path = _default_file_path(label)
            if default_path.exists():
                file_paths[label] = str(default_path)

    if "tb_current" not in file_paths or "tb_previous" not in file_paths:
        return None

    uploaded_files_store[session_id] = file_paths
    logger.info(f"Restored upload session {session_id} from disk")
    return file_paths


def _get_session_files(session_id: str) -> dict | None:
    return uploaded_files_store.get(session_id) or _restore_session_files(session_id)


def _validate_complete_session(session_id: str) -> dict:
    file_paths = _get_session_files(session_id)
    if not file_paths:
        raise HTTPException(
            status_code=400,
            detail=(
                "Upload session expired or files are not available on this server. "
                "Please upload the files again and generate the report."
            ),
        )

    missing_labels = [
        label
        for label in ["tb_current", "tb_previous", "budget", "mapping"]
        if label not in file_paths or not Path(file_paths[label]).exists()
    ]
    if missing_labels:
        raise HTTPException(
            status_code=400,
            detail=(
                "Upload session is incomplete. Missing files: "
                f"{', '.join(missing_labels)}. Please upload the files again."
            ),
        )

    return file_paths


def _build_report_for_session(
    session_id: str,
    db: Session,
    user_id: int,
    user_name: str,
) -> ReportResponse:
    start_time = time.time()
    file_paths = _validate_complete_session(session_id)
    logger.info(f"Starting report generation for session {session_id}")

    report_month, report_year = detect_tb_month_and_year(file_paths["tb_current"])
    if not report_month or not report_year:
        raise HTTPException(
            status_code=400,
            detail="Could not detect month and year from Trial Balance filename. "
            "Please ensure the filename contains a month abbreviation and year.",
        )
    logger.info(f"Report period: {report_month} {report_year}")

    logger.info("Loading mapping rules...")
    mapping_engine = MappingEngine()
    mapping_engine.load_mapping(file_paths["mapping"])

    logger.info("Reading current year Trial Balance...")
    cy_tb_df = read_current_year_tb(file_paths["tb_current"])
    if cy_tb_df.empty:
        raise HTTPException(
            status_code=400,
            detail="Current Year Trial Balance contains no valid data rows.",
        )

    logger.info("Mapping current year TB rows...")
    cy_tb_df = mapping_engine.map_tb_rows(cy_tb_df)
    mapped_count = (cy_tb_df["revenue_category"] != "").sum()
    unmapped_count = (cy_tb_df["revenue_category"] == "").sum()
    cy_unmapped_df = mapping_engine.get_unmapped_report()
    logger.info(f"CY TB: {mapped_count} mapped, {unmapped_count} unmapped")

    logger.info("Reading previous year Trial Balance...")
    py_tb_df = read_previous_year_tb(file_paths["tb_previous"])
    if py_tb_df.empty:
        logger.warning("Previous Year TB is empty, proceeding without PY data")
        import pandas as pd

        py_tb_df = pd.DataFrame(columns=cy_tb_df.columns)

    logger.info("Mapping previous year TB rows...")
    py_tb_df = mapping_engine.map_tb_rows(py_tb_df)
    py_unmapped_df = mapping_engine.get_unmapped_report()

    file_paths["cy_unmapped_df"] = cy_unmapped_df
    file_paths["py_unmapped_df"] = py_unmapped_df
    file_paths["report_month"] = report_month
    file_paths["report_year"] = report_year

    session_dir = settings.UPLOAD_DIR / session_id
    if session_dir.exists():
        try:
            cy_unmapped_df.to_csv(session_dir / "cy_unmapped.csv", index=False)
            py_unmapped_df.to_csv(session_dir / "py_unmapped.csv", index=False)
            (session_dir / "report_meta.json").write_text(
                json.dumps({"report_month": report_month, "report_year": report_year}),
                encoding="utf-8"
            )
        except Exception as e:
            logger.warning(f"Failed to persist unmapped DataFrames to disk: {e}")

    logger.info("Processing current year TB data...")
    cy_tb_data = process_trial_balance(cy_tb_df)

    logger.info("Processing previous year TB data...")
    py_tb_data = process_trial_balance(py_tb_df)

    logger.info("Loading budget data...")
    budget_proc = BudgetProcessor()
    budget_proc.load_budget(file_paths["budget"], report_month, report_year)

    logger.info("Calculating revenue...")
    revenue_results = calculate_revenue(cy_tb_data, py_tb_data, budget_proc)

    logger.info("Generating PowerPoint presentation...")
    ppt_path = generate_pptx(
        revenue_results,
        report_month,
        report_year,
        str(settings.OUTPUT_DIR),
    )

    elapsed = round(time.time() - start_time, 2)
    filename = Path(ppt_path).name

    logger.info(f"Report generated successfully in {elapsed}s: {filename}")

    log_audit(
        db=db,
        action="Report Generation",
        module="Finance",
        description=f"Generated PowerPoint report. Filename: {filename}, Month: {report_month}, Year: {report_year}, Mapped: {mapped_count}, Unmapped: {unmapped_count}",
        user_id=user_id,
        user_name=user_name,
    )

    return ReportResponse(
        status="success",
        filename=filename,
        unmapped_count=unmapped_count,
        total_mapped=mapped_count,
        processing_time_seconds=elapsed,
        report_month=report_month,
        report_year=report_year,
        message=f"Report generated for {report_month} {report_year}. "
        f"{mapped_count} records mapped, {unmapped_count} unmapped.",
    )


def _run_report_job(session_id: str, user_id: int, user_name: str):
    db = SessionLocal()
    try:
        result = _build_report_for_session(session_id, db, user_id, user_name)
        job_data = result.model_dump()
        report_jobs[session_id] = job_data
        # Persist to DB so other workers can read the full result
        _db_save_job(session_id, "success", job_data)
    except HTTPException as e:
        detail = e.detail
        if isinstance(detail, dict) and "errors" in detail:
            message = ", ".join(detail["errors"])
        else:
            message = str(detail)
        logger.error(f"Report job failed for session {session_id}: {message}")
        error_data = {"status": "error", "message": message}
        report_jobs[session_id] = error_data
        _db_save_job(session_id, "error", error_data)
    except Exception as e:
        logger.exception(f"Report job failed for session {session_id}: {e}")
        error_data = {
            "status": "error",
            "message": f"Report generation failed: {str(e)}",
        }
        report_jobs[session_id] = error_data
        _db_save_job(session_id, "error", error_data)
    finally:
        db.close()


@router.post("/upload")
async def upload_files(
    request: Request,
    tb_current: UploadFile = File(...),
    tb_previous: UploadFile = File(...),
    budget: UploadFile | None = File(None),
    mapping: UploadFile | None = File(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    logger.info(f"User {current_user.email} is uploading files")

    session_id = str(uuid.uuid4())[:8]
    session_dir = settings.UPLOAD_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    file_paths = {}
    files = {
        "tb_current": tb_current,
        "tb_previous": tb_previous,
    }
    if budget is not None:
        files["budget"] = budget
    if mapping is not None:
        files["mapping"] = mapping

    for label, upload_file in files.items():
        ext = Path(upload_file.filename).suffix.lower()

    # Allow TXT only for Trial Balance files
        if label in ["tb_current", "tb_previous"]:
            allowed_extensions = {".xlsx", ".xls", ".txt"}
        else:
            allowed_extensions = {".xlsx", ".xls"}

        if ext not in allowed_extensions:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid file type for {label}: {ext}. Allowed: {', '.join(sorted(allowed_extensions))}",
        )

        save_path = session_dir / f"{label}__{upload_file.filename}"

        with open(save_path, "wb") as f:
            content = await upload_file.read()
            f.write(content)

        file_paths[label] = str(save_path)
        logger.info(f"Saved {label}: {upload_file.filename} -> {save_path}")

    for label in ["budget", "mapping"]:
        if label not in file_paths:
            default_path = _default_file_path(label)
            if default_path.exists():
                file_paths[label] = str(default_path)
                logger.info(f"Using default {label} workbook: {default_path}")

    validation = validate_uploaded_files(
        file_paths.get("tb_current"),
        file_paths.get("tb_previous"),
        file_paths.get("budget"),
        file_paths.get("mapping"),
    )

    if not validation.is_valid:
        raise HTTPException(status_code=400, detail={"errors": validation.errors})

    tb_val = validate_trial_balance(file_paths["tb_current"], "Current Year")
    if not tb_val.is_valid:
        raise HTTPException(status_code=400, detail={"errors": tb_val.errors})

    py_val = validate_trial_balance(file_paths["tb_previous"], "Previous Year")
    if not py_val.is_valid:
        raise HTTPException(status_code=400, detail={"errors": py_val.errors})

    bud_val = validate_budget_workbook(file_paths["budget"])
    if not bud_val.is_valid:
        raise HTTPException(status_code=400, detail={"errors": bud_val.errors})

    map_val = validate_mapping_workbook(file_paths["mapping"])
    if not map_val.is_valid:
        raise HTTPException(status_code=400, detail={"errors": map_val.errors})

    uploaded_files_store[session_id] = file_paths

    log_audit(
        db=db,
        action="File Upload",
        module="Finance",
        description=f"Uploaded files for processing. Session ID: {session_id}, Files: TB Current={tb_current.filename}, TB Previous={tb_previous.filename}",
        user_id=current_user.id,
        user_name=current_user.full_name,
        request=request
    )

    return {
        "status": "success",
        "session_id": session_id,
        "files": {
            "tb_current": tb_current.filename,
            "tb_previous": tb_previous.filename,
            "budget": budget.filename if budget else Path(file_paths["budget"]).name,
            "mapping": mapping.filename if mapping else Path(file_paths["mapping"]).name,
        },
        "warnings": validation.warnings,
    }


@router.get("/admin/config")
async def get_admin_config(current_user: User = Depends(get_current_active_user)):
    return {
        **_default_file_info("budget"),
        **_default_file_info("mapping"),
    }


@router.post("/admin/upload-default")
async def upload_default_file(
    request: Request,
    file_type: str = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: User = Depends(get_admin_user),
):
    if file_type not in {"budget", "mapping"}:
        raise HTTPException(status_code=400, detail="file_type must be 'budget' or 'mapping'.")

    ext = Path(file.filename).suffix.lower()
    if ext not in {".xlsx", ".xls"}:
        raise HTTPException(status_code=400, detail="Default files must be Excel workbooks.")

    target_path = DEFAULT_FILES_DIR / f"default_{file_type}{ext}"
    for old_path in DEFAULT_FILES_DIR.glob(f"default_{file_type}.*"):
        old_path.unlink(missing_ok=True)

    with open(target_path, "wb") as f:
        content = await file.read()
        f.write(content)

    validation = (
        validate_budget_workbook(str(target_path))
        if file_type == "budget"
        else validate_mapping_workbook(str(target_path))
    )
    if not validation.is_valid:
        target_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail={"errors": validation.errors})

    log_audit(
        db=db,
        action="Default File Upload",
        module="Finance Admin",
        description=f"Uploaded default {file_type} workbook: {file.filename}",
        user_id=admin.id,
        user_name=admin.full_name,
        request=request,
    )

    return {
        "status": "success",
        "file_type": file_type,
        "filename": target_path.name,
        "warnings": validation.warnings,
    }


@router.post("/generate")
async def generate_report(
    background_tasks: BackgroundTasks,
    session_id: str = "",
    current_user: User = Depends(get_current_active_user)
):
    if not session_id:
        raise HTTPException(
            status_code=400, detail="Invalid or missing session_id. Upload files first."
        )

    _validate_complete_session(session_id)
    existing_job = report_jobs.get(session_id)
    if existing_job:
        return existing_job

    # Mark as processing in DB (visible to other workers) and launch background task
    _db_save_job(session_id, "processing")
    report_jobs[session_id] = {
        "status": "processing",
        "message": "Report generation started.",
    }
    background_tasks.add_task(
        _run_report_job,
        session_id,
        current_user.id,
        current_user.full_name,
    )

    return report_jobs[session_id]


@router.get("/report-status/{session_id}")
async def get_report_status(
    session_id: str,
    current_user: User = Depends(get_current_active_user)
):
    # 1. Check in-memory (same worker — fastest path)
    job = report_jobs.get(session_id)
    if job and job.get("status") in ("success", "error"):
        return job

    # 2. Query the database (works across all workers and restarts)
    db_result = _db_load_job(session_id)
    if db_result:
        if db_result.get("status") in ("success", "error"):
            # Cache in this worker's memory for future polls
            report_jobs[session_id] = db_result
        return db_result

    # 3. If in-memory shows "processing" but DB has nothing yet, keep polling
    if job:
        return job

    # 4. Truly unknown session — tell client to keep waiting (do NOT return fake success)
    return {"status": "processing", "message": "Waiting for report generation to start."}


@router.get("/download/{filename}")
async def download_report(
    filename: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    file_path = settings.OUTPUT_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Report file not found.")

    log_audit(
        db=db,
        action="Download Report",
        module="Finance",
        description=f"Downloaded report PowerPoint file: {filename}",
        user_id=current_user.id,
        user_name=current_user.full_name,
        request=request
    )

    return FileResponse(
        path=str(file_path),
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


def _get_or_compute_unmapped_dfs(session_id: str, file_paths: dict):
    import pandas as pd
    session_dir = settings.UPLOAD_DIR / session_id
    cy_unmapped = file_paths.get("cy_unmapped_df")
    py_unmapped = file_paths.get("py_unmapped_df")
    report_month = file_paths.get("report_month", "")
    report_year = file_paths.get("report_year", "")

    meta_file = session_dir / "report_meta.json"
    cy_csv = session_dir / "cy_unmapped.csv"
    py_csv = session_dir / "py_unmapped.csv"

    if (cy_unmapped is None or py_unmapped is None) and cy_csv.exists() and py_csv.exists():
        try:
            cy_unmapped = pd.read_csv(cy_csv).fillna("")
            py_unmapped = pd.read_csv(py_csv).fillna("")
            file_paths["cy_unmapped_df"] = cy_unmapped
            file_paths["py_unmapped_df"] = py_unmapped
        except Exception as e:
            logger.warning(f"Error loading unmapped CSVs for session {session_id}: {e}")

    if (not report_month or not report_year) and meta_file.exists():
        try:
            meta = json.loads(meta_file.read_text(encoding="utf-8"))
            report_month = meta.get("report_month", "")
            report_year = meta.get("report_year", "")
            file_paths["report_month"] = report_month
            file_paths["report_year"] = report_year
        except Exception:
            pass

    if cy_unmapped is None or py_unmapped is None or not report_month or not report_year:
        logger.info(f"Re-generating unmapped DataFrames on the fly for session {session_id}")
        tb_current_path = file_paths.get("tb_current")
        tb_previous_path = file_paths.get("tb_previous")
        mapping_path = file_paths.get("mapping")

        if not tb_current_path or not tb_previous_path:
            return None, None, "", ""

        report_month, report_year = detect_tb_month_and_year(tb_current_path)

        mapping_engine = MappingEngine()
        if mapping_path and os.path.exists(mapping_path):
            mapping_engine.load_mapping(mapping_path)

        cy_tb_df = read_current_year_tb(tb_current_path)
        if not cy_tb_df.empty:
            cy_tb_df = mapping_engine.map_tb_rows(cy_tb_df)
            cy_unmapped = mapping_engine.get_unmapped_report()
        else:
            cy_unmapped = pd.DataFrame()

        py_tb_df = read_previous_year_tb(tb_previous_path)
        if not py_tb_df.empty:
            py_tb_df = mapping_engine.map_tb_rows(py_tb_df)
            py_unmapped = mapping_engine.get_unmapped_report()
        else:
            py_unmapped = pd.DataFrame()

        file_paths["cy_unmapped_df"] = cy_unmapped
        file_paths["py_unmapped_df"] = py_unmapped
        file_paths["report_month"] = report_month
        file_paths["report_year"] = report_year

        if session_dir.exists():
            try:
                cy_unmapped.to_csv(cy_csv, index=False)
                py_unmapped.to_csv(py_csv, index=False)
                meta_file.write_text(json.dumps({"report_month": report_month, "report_year": report_year}), encoding="utf-8")
            except Exception:
                pass

    return cy_unmapped, py_unmapped, report_month, report_year


def _write_summary_and_detail_sheets(
    writer, cy_unmapped, py_unmapped, report_month, report_year
):
    import pandas as pd
    import re as _re
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    cy_count = len(cy_unmapped) if not cy_unmapped.empty else 0
    py_count = len(py_unmapped) if not py_unmapped.empty else 0

    def _num(df, col):
        return pd.to_numeric(df.get(col, 0), errors="coerce").fillna(0)

    def _fmt_pa(val):
        if val == 0:
            return "0"
        sign = "+" if val > 0 else ""
        return f"{sign}{val:,.0f}"

    def _categorize_reason(reason):
        if reason.startswith("International filter:"):
            m = _re.search(r"BL (\d+)", reason)
            bl = m.group(1) if m else "?"
            return "International BL Filter", bl, reason
        if reason.startswith("Equipment Sales filter:"):
            m = _re.search(r"Account (\d+)", reason)
            acct = m.group(1) if m else "?"
            return "Equipment Sales Cleanup", acct, reason
        if reason.startswith("GL ") and "not covered" in reason:
            m = _re.search(r"GL (\d+)", reason)
            acct = m.group(1) if m else "?"
            return "GL Account Not Covered", acct, reason
        if "Product" in reason and "outside" in reason:
            m = _re.search(r"Product (\d+).*GL (\d+)", reason)
            key = f"P{m.group(1)}_GL{m.group(2)}" if m else reason[:30]
            return "Product Range Gap", key, reason
        if "BL" in reason and "outside" in reason:
            m = _re.search(r"BL (\d+)", reason)
            bl = m.group(1) if m else "?"
            return "BL Range Gap", bl, reason
        return "Other", reason[:40], reason

    def _get_category_info(issue_name: str, issue_key: str, sample_reason: str, sample_row: dict = None):
        gl_code = ""
        product_code = ""
        bl_code = ""
        if sample_row:
            gl_code = str(sample_row.get("gl_code") or "").strip()
            product_code = str(sample_row.get("product") or "").strip()
            bl_code = str(sample_row.get("business_line") or "").strip()

        if not gl_code and issue_key and issue_key.isdigit() and len(issue_key) >= 6:
            gl_code = issue_key

        if issue_name == "GL Account Not Covered":
            gl_str = f"[{gl_code}]" if gl_code else "[GL_Code]"
            desc = "This General Ledger account recorded revenue activity in the Trial Balance but is completely missing from your master mapping rules."
            action = f"Open 'Revenue Mapping Workbook' -> Navigate to 'Code Mapping' sheet -> Append GL Code {gl_str} and assign its appropriate Revenue Category."
            target_seg = "account"
        elif issue_name == "Product Range Gap":
            gl_str = f"[{gl_code}]" if gl_code else "[GL_Code]"
            prod_str = f"[{product_code}]" if product_code else "[Product_Code]"
            desc = f"The GL account is mapped, but Product Segment {prod_str} falls outside the mapped Product Range defined for this account."
            action = f"Open 'Revenue Mapping Workbook' -> Locate GL Code {gl_str} -> Extend the Start/End Product Range to include Product {prod_str}."
            target_seg = "product"
        elif issue_name == "BL Range Gap":
            gl_str = f"[{gl_code}]" if gl_code else "[GL_Code]"
            bl_str = f"[{bl_code}]" if bl_code else "[BL_Code]"
            desc = f"Business Line Segment {bl_str} is not covered under the mapped ranges for this GL account."
            action = f"Open 'Revenue Mapping Workbook' -> Locate GL Code {gl_str} -> Update Business Line range to cover code {bl_str}."
            target_seg = "business_line"
        elif issue_name in ["International BL Filter", "Equipment Sales Cleanup"] or "Filter" in issue_name:
            bl_str = f"[{bl_code}]" if bl_code else "[BL_Code]"
            gl_str = f"[{gl_code}]" if gl_code else "[GL_Code]"
            desc = "The row hit a domain-specific filter rule (e.g., Equipment Sales/International) and was bypassed by standard category logic."
            action = f"Review domain filter rules in 'Revenue Mapping Workbook' -> Confirm if Business Line {bl_str} or Account {gl_str} requires dedicated override classification."
            target_seg = "business_line" if issue_name == "International BL Filter" else "account"
        else:
            desc = "Unmapped combination of GL Account, Product, and Business Line segments in the Code Mapping workbook."
            action = "Open 'Revenue Mapping Workbook' -> Check 'Code Mapping' sheet -> Ensure GL, Product, and Business Line segments cover this row."
            target_seg = "account"

        return {
            "finance_explanation": desc,
            "how_to_fix": action,
            "target_segment": target_seg,
        }

    _DARK_BLUE = "1F3864"
    _MED_BLUE = "2E75B6"
    _LIGHT_BLUE = "D6E4F0"
    _WHITE = "FFFFFF"
    _LIGHT_GRAY = "F2F2F2"
    _RED = "C00000"
    _AMBER_FILL = "FFF2CC"
    _AMBER_TEXT = "7F6000"

    hdr_font = Font(name="Calibri", bold=True, color=_WHITE, size=11)
    hdr_fill = PatternFill(start_color=_MED_BLUE, end_color=_MED_BLUE, fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_font = Font(name="Calibri", bold=True, color=_WHITE, size=14)
    title_fill = PatternFill(start_color=_DARK_BLUE, end_color=_DARK_BLUE, fill_type="solid")

    subtitle_font = Font(name="Calibri", bold=True, color=_DARK_BLUE, size=11)

    total_font = Font(name="Calibri", bold=True, size=11)
    total_fill = PatternFill(start_color=_LIGHT_BLUE, end_color=_LIGHT_BLUE, fill_type="solid")

    data_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill(start_color=_LIGHT_GRAY, end_color=_LIGHT_GRAY, fill_type="solid")

    amber_fill = PatternFill(start_color=_AMBER_FILL, end_color=_AMBER_FILL, fill_type="solid")
    amber_font = Font(name="Calibri", bold=True, color=_AMBER_TEXT, size=10)

    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    def _auto_width(ws, min_width=10, max_width=65):
        for col_cells in ws.columns:
            col_letter = get_column_letter(col_cells[0].column)
            lengths = []
            for cell in col_cells:
                if cell.value is not None:
                    lengths.append(len(str(cell.value)))
            if lengths:
                w = min(max(max(lengths) + 2, min_width), max_width)
                ws.column_dimensions[col_letter].width = w

    def _style_header_row(ws, row_num, num_cols):
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = hdr_align
            cell.border = thin_border

    def _style_data_rows(ws, start_row, end_row, num_cols):
        for r in range(start_row, end_row + 1):
            is_alt = (r - start_row) % 2 == 1
            for c in range(1, num_cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if is_alt and cell.fill.fill_type != "solid":
                    cell.fill = alt_fill

    def _style_total_row(ws, row_num, num_cols):
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=c)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thin_border

    if cy_unmapped.empty and py_unmapped.empty:
        ws = writer.book.create_sheet("SUMMARY")
        ws.append([f"Unmapped Rows Analysis - {report_month} {report_year}"])
        ws.append(["No unmapped rows found."])
        return

    all_parts = []
    if not cy_unmapped.empty:
        cy = cy_unmapped.copy()
        cy["_source"] = "CY"
        all_parts.append(cy)
    if not py_unmapped.empty:
        py = py_unmapped.copy()
        py["_source"] = "PY"
        all_parts.append(py)
    combined = pd.concat(all_parts, ignore_index=True)
    combined["period_activity"] = _num(combined, "period_activity")
    combined["ending_balance"] = _num(combined, "ending_balance")

    combined["_issue_name"] = combined["unmapped_reason"].apply(
        lambda r: _categorize_reason(r)[0]
    )
    combined["_issue_key"] = combined["unmapped_reason"].apply(
        lambda r: _categorize_reason(r)[1]
    )
    combined["_issue_reason"] = combined["unmapped_reason"].apply(
        lambda r: _categorize_reason(r)[2]
    )

    issue_types = (
        combined.groupby("_issue_name")
        .agg(
            CY_PA=("period_activity", "sum"),
            CY_Rows=("_source", lambda x: (x == "CY").sum()),
            PY_Rows=("_source", lambda x: (x == "PY").sum()),
            Sample_Reason=("_issue_reason", "first"),
            Sample_Flexfield=("flexfield", "first"),
            Sample_GL=("gl_code", "first"),
            Sample_Product=("product", "first"),
            Sample_BL=("business_line", "first"),
        )
        .reset_index()
    )
    issue_types["abs_CY_PA"] = issue_types["CY_PA"].abs()
    issue_types = issue_types.sort_values("abs_CY_PA", ascending=False).drop(
        columns=["abs_CY_PA"]
    )
    issue_types = issue_types.reset_index(drop=True)
    issue_types.index = issue_types.index + 1
    issue_types.index.name = "Issue #"

    priority_map = {
        "International BL Filter": "HIGH",
        "Equipment Sales Cleanup": "HIGH",
        "GL Account Not Covered": "MEDIUM",
        "Product Range Gap": "MEDIUM",
        "BL Range Gap": "LOW",
        "Other": "LOW",
    }

    # ----------------------------------------------------
    # SUMMARY SHEET GENERATION
    # ----------------------------------------------------
    ws_summary = writer.book.create_sheet("SUMMARY", 0)

    ws_summary.append([f"Unmapped Rows Finance Action Guide - {report_month} {report_year}"])
    ws_summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    title_cell = ws_summary.cell(row=1, column=1)
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_summary.append([
        f"CY Total Unmapped: {cy_count} rows ({_fmt_pa(combined[combined['_source']=='CY']['period_activity'].sum())} Rs. PA) | "
        f"PY Total Unmapped: {py_count} rows | Follow step-by-step instructions below to update Revenue Mapping Workbook."
    ])
    ws_summary.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    sub_cell = ws_summary.cell(row=2, column=1)
    sub_cell.font = subtitle_font
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws_summary.append([])

    summary_headers = [
        "Issue #",
        "Unmapped Category",
        "Detail Sheet",
        "CY Rows",
        "CY PA (Rs.)",
        "PY Rows",
        "Finance Description",
        "How to Fix (Workbook Guide)",
        "Priority",
    ]
    ws_summary.append(summary_headers)
    _style_header_row(ws_summary, 4, len(summary_headers))

    total_cy_pa = 0
    total_cy_rows = 0
    total_py_rows = 0
    data_start = 5
    for idx, row in issue_types.iterrows():
        priority = priority_map.get(row["_issue_name"], "LOW")
        sheet_label = f"{idx}. {row['_issue_name']}"
        sample_dict = {
            "gl_code": row.get("Sample_GL", ""),
            "product": row.get("Sample_Product", ""),
            "business_line": row.get("Sample_BL", ""),
        }
        info = _get_category_info(
            row["_issue_name"], row["_issue_name"], row["Sample_Reason"], sample_dict
        )

        ws_summary.append([
            idx,
            row["_issue_name"],
            sheet_label,
            int(row["CY_Rows"]),
            _fmt_pa(row["CY_PA"]),
            int(row["PY_Rows"]),
            info["finance_explanation"],
            info["how_to_fix"],
            priority,
        ])
        total_cy_pa += row["CY_PA"]
        total_cy_rows += int(row["CY_Rows"])
        total_py_rows += int(row["PY_Rows"])

    data_end = data_start + len(issue_types) - 1
    _style_data_rows(ws_summary, data_start, data_end, len(summary_headers))

    for r in range(data_start, data_end + 1):
        prio_cell = ws_summary.cell(row=r, column=9)
        if prio_cell.value == "HIGH":
            prio_cell.font = Font(name="Calibri", bold=True, color=_RED, size=10)

    ws_summary.append([])
    total_row_num = data_end + 2
    ws_summary.append([
        "TOTAL",
        "",
        "",
        total_cy_rows,
        _fmt_pa(total_cy_pa),
        total_py_rows,
        "",
        "",
        "",
    ])
    _style_total_row(ws_summary, total_row_num, len(summary_headers))

    ws_summary.append([
        f"Current Generated Total Revenue YTD includes these unmapped amounts. "
        f"Total unmapped CY PA: {_fmt_pa(total_cy_pa)} Rs."
    ])
    _auto_width(ws_summary)

    ws_summary.column_dimensions["A"].width = 10
    ws_summary.column_dimensions["B"].width = 28
    ws_summary.column_dimensions["C"].width = 28
    ws_summary.column_dimensions["D"].width = 12
    ws_summary.column_dimensions["E"].width = 18
    ws_summary.column_dimensions["F"].width = 12
    ws_summary.column_dimensions["G"].width = 50
    ws_summary.column_dimensions["H"].width = 65
    ws_summary.column_dimensions["I"].width = 12

    # ----------------------------------------------------
    # DETAIL SHEETS GENERATION (FAST SINGLE PASS)
    # ----------------------------------------------------
    detail_cols = [
        "Source",
        "GL Code",
        "Description",
        "Flexfield",
        "Cost Center",
        "Location",
        "Business Line",
        "Product",
        "Account",
        "Technology",
        "Intercompany",
        "Project",
        "Beginning Balance",
        "Period Activity",
        "Ending Balance",
        "Suggested Fix / Required Action",
        "Unmapped Reason",
    ]

    align_center = Alignment(vertical="center", wrap_text=True)

    for issue_idx, (issue_name, type_df) in enumerate(
        combined.groupby("_issue_name"), start=1
    ):
        cy_type = type_df[type_df["_source"] == "CY"]
        py_type = type_df[type_df["_source"] == "PY"]
        cy_type_pa = cy_type["period_activity"].sum()
        cy_type_rows = len(cy_type)
        py_type_rows = len(py_type)

        records = type_df.to_dict("records")
        sample_row_dict = records[0] if records else {}
        cat_info = _get_category_info(issue_name, issue_name, sample_row_dict)
        target_seg = cat_info["target_segment"]

        sheet_label = f"{issue_idx}. {issue_name}"
        ws = writer.book.create_sheet(sheet_label)

        # Header Callout Banner Block
        ws.append([f"ISSUE RESOLUTION GUIDE: {issue_idx}. {issue_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=17)
        c1 = ws.cell(row=1, column=1)
        c1.font = title_font
        c1.fill = title_fill
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.append([f"EXPLANATION: {cat_info['finance_explanation']}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=17)
        c2 = ws.cell(row=2, column=1)
        c2.font = Font(name="Calibri", bold=True, color=_DARK_BLUE, size=10)
        c2.fill = PatternFill(start_color=_LIGHT_BLUE, end_color=_LIGHT_BLUE, fill_type="solid")
        c2.alignment = Alignment(horizontal="left", vertical="center")

        ws.append([f"ACTION REQUIRED: {cat_info['how_to_fix']}"])
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=17)
        c3 = ws.cell(row=3, column=1)
        c3.font = Font(name="Calibri", bold=True, color="7F6000", size=10)
        c3.fill = amber_fill

        ws.append([
            f"IMPACT METRICS: CY Total: {cy_type_rows} rows ({_fmt_pa(cy_type_pa)} Rs. PA) | "
            f"PY Total: {py_type_rows} rows | Amber highlighted segment columns indicate the unmapped criteria."
        ])
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=17)
        c4 = ws.cell(row=4, column=1)
        c4.font = Font(name="Calibri", italic=True, size=10)
        c4.fill = alt_fill

        ws.append([])
        ws.append(detail_cols)
        hdr_r = ws.max_row
        for c in range(1, len(detail_cols) + 1):
            cell = ws.cell(row=hdr_r, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = hdr_align
            cell.border = thin_border

        for row_i, r in enumerate(records):
            gl_val = str(r.get("gl_code") or "").strip()
            prod_val = str(r.get("product") or "").strip()
            bl_val = str(r.get("business_line") or "").strip()

            if issue_name == "GL Account Not Covered":
                row_fix = f"Open 'Revenue Mapping Workbook' -> 'Code Mapping' sheet -> Add GL Account {gl_val} and assign Revenue Category"
            elif issue_name == "Product Range Gap":
                row_fix = f"Open 'Revenue Mapping Workbook' -> 'Code Mapping' sheet -> Extend Product Range for GL Account {gl_val} to include Product {prod_val}"
            elif issue_name == "BL Range Gap":
                row_fix = f"Open 'Revenue Mapping Workbook' -> 'Code Mapping' sheet -> Update Business Line Range for GL Account {gl_val} to include BL {bl_val}"
            elif issue_name == "International BL Filter":
                row_fix = f"Review International BL rule in 'Revenue Mapping Workbook' for Business Line {bl_val} / Account {gl_val}"
            elif issue_name == "Equipment Sales Cleanup":
                row_fix = f"Review Equipment Sales account rule in 'Revenue Mapping Workbook' for Account {gl_val}"
            else:
                row_fix = f"Open 'Revenue Mapping Workbook' -> 'Code Mapping' sheet -> Update rule for GL {gl_val}, Product {prod_val}, BL {bl_val}"

            row_data = [
                r.get("_source", ""),
                r.get("gl_code", ""),
                r.get("description", ""),
                r.get("flexfield", ""),
                r.get("cost_center", ""),
                r.get("location", ""),
                r.get("business_line", ""),
                r.get("product", ""),
                r.get("account", ""),
                r.get("technology", ""),
                r.get("intercompany", ""),
                r.get("project", ""),
                _fmt_pa(r.get("beginning_balance", 0)),
                _fmt_pa(r.get("period_activity", 0)),
                _fmt_pa(r.get("ending_balance", 0)),
                row_fix,
                r.get("unmapped_reason", ""),
            ]

            ws.append(row_data)
            curr_r = ws.max_row
            is_alt = row_i % 2 == 1

            # Single-pass cell styling during row insertion
            for col_i in range(1, len(detail_cols) + 1):
                cell = ws.cell(row=curr_r, column=col_i)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = align_center

                if target_seg == "account" and col_i in (2, 9):
                    cell.fill = amber_fill
                    cell.font = amber_font
                elif target_seg == "product" and col_i == 8:
                    cell.fill = amber_fill
                    cell.font = amber_font
                elif target_seg == "business_line" and col_i == 7:
                    cell.fill = amber_fill
                    cell.font = amber_font
                elif is_alt:
                    cell.fill = alt_fill

        ws.append([])
        ws.append([
            "TOTAL",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            _fmt_pa(type_df["beginning_balance"].sum() if "beginning_balance" in type_df.columns else 0),
            _fmt_pa(cy_type_pa),
            _fmt_pa(type_df["ending_balance"].sum() if "ending_balance" in type_df.columns else 0),
            "",
            "",
        ])
        tot_r = ws.max_row
        for c in range(1, len(detail_cols) + 1):
            cell = ws.cell(row=tot_r, column=c)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thin_border

        # Set column dimensions directly
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 40
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 10
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 12
        ws.column_dimensions["I"].width = 14
        ws.column_dimensions["J"].width = 12
        ws.column_dimensions["K"].width = 14
        ws.column_dimensions["L"].width = 10
        ws.column_dimensions["M"].width = 18
        ws.column_dimensions["N"].width = 18
        ws.column_dimensions["O"].width = 18
        ws.column_dimensions["P"].width = 75
        ws.column_dimensions["Q"].width = 50


@router.post("/generate-unmapped")
async def generate_unmapped_report(
    request: Request,
    session_id: str = Query(""),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    if not session_id:
        raise HTTPException(
            status_code=400,
            detail="Invalid or missing session_id. Generate report first.",
        )

    file_paths = _get_session_files(session_id)
    if not file_paths:
        raise HTTPException(
            status_code=400,
            detail=(
                "Upload session expired or files are not available on this server. "
                "Please upload the files and generate the report again."
            ),
        )

    cy_unmapped, py_unmapped, report_month, report_year = _get_or_compute_unmapped_dfs(session_id, file_paths)

    if cy_unmapped is None or py_unmapped is None:
        raise HTTPException(
            status_code=400,
            detail="Unmapped data not available. Please generate the report first.",
        )

    try:
        import pandas as pd

        output_filename = f"Unmapped_Rows_{report_month}_{report_year}.xlsx"
        output_path = settings.OUTPUT_DIR / output_filename

        display_cols = [
            "gl_code",
            "description",
            "flexfield",
            "cost_center",
            "location",
            "business_line",
            "product",
            "account",
            "technology",
            "intercompany",
            "project",
            "beginning_balance",
            "period_activity",
            "ending_balance",
        ]

        with pd.ExcelWriter(str(output_path), engine="openpyxl") as writer:
            if not cy_unmapped.empty:
                cy_export = cy_unmapped[
                    [c for c in display_cols if c in cy_unmapped.columns]
                ].copy()
                cy_export.to_excel(writer, sheet_name="CY Unmapped", index=False)
            else:
                pd.DataFrame(columns=display_cols).to_excel(
                    writer, sheet_name="CY Unmapped", index=False
                )

            if not py_unmapped.empty:
                py_export = py_unmapped[
                    [c for c in display_cols if c in py_unmapped.columns]
                ].copy()
                py_export.to_excel(writer, sheet_name="PY Unmapped", index=False)
            else:
                pd.DataFrame(columns=display_cols).to_excel(
                    writer, sheet_name="PY Unmapped", index=False
                )

            _write_summary_and_detail_sheets(
                writer, cy_unmapped, py_unmapped, report_month, report_year
            )

        cy_count = len(cy_unmapped) if not cy_unmapped.empty else 0
        py_count = len(py_unmapped) if not py_unmapped.empty else 0
        logger.info(
            f"Unmapped report generated: {output_filename} (CY: {cy_count}, PY: {py_count})"
        )

        log_audit(
            db=db,
            action="Unmapped Report Generation",
            module="Finance",
            description=f"Generated unmapped rows Excel report. Filename: {output_filename}, Month: {report_month}, Year: {report_year}, CY: {cy_count}, PY: {py_count}",
            user_id=current_user.id,
            user_name=current_user.full_name,
            request=request
        )

        return {
            "status": "success",
            "filename": output_filename,
            "cy_unmapped_count": cy_count,
            "py_unmapped_count": py_count,
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"Error generating unmapped report: {e}")
        raise HTTPException(
            status_code=500, detail=f"Unmapped report generation failed: {str(e)}"
        )


@router.get("/download-unmapped/{filename}")
async def download_unmapped_report(
    filename: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    file_path = settings.OUTPUT_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Unmapped report file not found.")

    log_audit(
        db=db,
        action="Download Unmapped Rows",
        module="Finance",
        description=f"Downloaded unmapped rows Excel report: {filename}",
        user_id=current_user.id,
        user_name=current_user.full_name,
        request=request
    )

    return FileResponse(
        path=str(file_path),
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/status")
async def health_check():
    return {"status": "ok", "app": settings.APP_NAME, "version": settings.APP_VERSION}
    return FileResponse(
        path=str(file_path),
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/status")
async def health_check():
    return {"status": "ok", "app": settings.APP_NAME, "version": settings.APP_VERSION}
