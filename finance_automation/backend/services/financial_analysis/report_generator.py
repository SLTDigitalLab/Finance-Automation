"""
Writes the quarterly Financial Analysis Report (Word) in Finance's exact format.

The template is Finance's own quarterly report (matching Output file_Q2 2026 Analysis.docx).
The generator preserves the entire layout, formatting, tables, and structure of the reference
document:
  - Executive Summary (SLT Group, SLT PLC populated from TB, Mobitel)
  - Financial Statistics – SLT PLC (Financial Summary and Key Ratios updated with actuals)
  - Operational Statistics – SLT PLC (Key Drivers & Capacity vs Utilization tables preserved)
  - Section (A) Quarter vs same quarter last year (15-col Group table with SLT PLC populated,
    5-col SLT PLC table populated, charts & commentary, Mobitel commentary)
  - Section (B) Quarter vs preceding quarter (15-col Group table, 5-col SLT PLC table, commentary, Mobitel)
  - Section (C) YTD vs preceding year's YTD (15-col Group table, 5-col SLT PLC table, commentary,
    Liquidity & Cashflow table, Borrowings table, Interest cost table, Mobitel commentary)
  - Section (D) Actuals vs Budget (SLT PLC table with actuals, budget note & commentary, Mobitel)
  - Key Projects Info – SLT PLC (FTTH Service-wise and RTOM-wise utilization tables preserved)
  - Annexure (i) – Borrowings Re-payment Schedule (preserved)
  - Annexure (ii) – Account-level variance drivers (GL drivers for Revenue and Opex across YoY, QoQ, YTD)
  - Sections to be completed by Finance & Data notes
"""
import calendar
import copy
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from services.financial_analysis import commentary
from services.financial_analysis.commentary import money, percent
from services.financial_analysis.pl_engine import Comparison, FinancialAnalysis, MonthFigures
from services.financial_analysis.periods import Period, Quarter
from utils.logger import logger

TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "financial_analysis_templates" / "financial_analysis_template.docx"
ENTITY = "SLT PLC"
DRIVER_ROWS = 8
FAVOURABLE, UNFAVOURABLE = "00B050", "EE0000"   # the template's green / red for growth %
DRIVER_COLUMNS = [3900, 1250, 1250, 1100, 1100]  # twips; a wider account column for the annexure

MANUAL_SECTIONS = [
    "SLT Group and Mobitel figures and commentary (the trial balance covers SLT PLC only)",
    "Financial Statistics items not in the P&L mapping: Gross Profit, EPS, Debt/Equity, ROE, Capex, "
    "free cash flows, staff cost to revenue and the balance sheet ratios",
    "Operational Statistics – SLT PLC (customer base, ARPU, capacity vs utilization)",
    "Liquidity and Operating Cashflow, Borrowings and Interest cost tables",
    "Key Projects Info – FTTH port utilization and FTTH active, suspended and churn customers",
    "Annexure (i) – Borrowings re-payment schedule",
    "Business reasons behind each movement (e.g. tariff revisions, salary increments)",
]


class ReportTemplateError(RuntimeError):
    """The Word template is missing a part the generator copies."""


@dataclass
class ReportInputs:
    analysis: FinancialAnalysis
    trend: List[MonthFigures]
    data_notes: List[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Labels


def _short(period: Period) -> str:
    return f"{period.label[:3]}’{period.year % 100:02d}"


def _long(period: Period, quote: str = "’") -> str:
    """'June’26' – the section (C) heading and the Financial Statistics tables spell the month out."""
    return f"{calendar.month_name[period.month]}{quote}{period.year % 100:02d}"


def _q_short(q: Quarter) -> str:
    return f"Q{q.quarter}’{q.year % 100:02d}"


def _ordinal(n: int) -> str:
    return {1: "st", 2: "nd", 3: "rd"}.get(n, "th")


@dataclass
class SectionSpec:
    letter: str
    title: str
    comparison_key: str
    current_label: str          # in sentences, e.g. "Q2’26"
    base_label: str
    period_phrase: str          # ends the revenue / opex sentences
    header_current: Sequence[str]   # table header lines
    header_base: Sequence[str]
    charts: bool = False


def _section_specs(q: Quarter) -> List[SectionSpec]:
    ly, prev = q.same_quarter_last_year(), q.previous()
    end, ly_end = q.end, ly.end
    return [
        SectionSpec(
            "A", f"Performance of the current quarter against the preceding year corresponding quarter "
                 f"({_q_short(q)} Vs {_q_short(ly)})",
            "quarter_vs_last_year", _q_short(q), _q_short(ly),
            f"in {_q_short(q)} compared to {_q_short(ly)}",
            [f"Q{q.quarter}", str(q.year)], [f"Q{ly.quarter}", str(ly.year)], charts=True),
        SectionSpec(
            "B", f"Performance of the current quarter against the preceding quarter "
                 f"({_q_short(q)} vs {_q_short(prev)})",
            "quarter_vs_previous", _q_short(q), _q_short(prev),
            f"in {_q_short(q)} compared to the preceding quarter ({_q_short(prev)})",
            [f"Q{q.quarter}", str(q.year)], [f"Q{prev.quarter}", str(prev.year)]),
        SectionSpec(
            "C", f"Performance of the current year against the preceding year "
                 f"(YTD {_short(end)} Vs YTD {_short(ly_end)})",
            "ytd_vs_last_year", f"YTD {_short(end)}", f"YTD {_short(ly_end)}",
            f"for YTD {_short(end)} compared to the corresponding period of the previous year",
            ["YTD", _short(end)], ["YTD", _short(ly_end)]),
    ]


# ---------------------------------------------------------------------------
# XML helpers


def _norm(text: str) -> str:
    return re.sub(r"[^a-z]", "", (text or "").lower())


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _set_text(p, text: str):
    """Replaces a paragraph's text, keeping its paragraph formatting and its first run's formatting."""
    runs = list(p.iter(qn("w:r")))
    keep = next((r for r in runs if r.find(qn("w:t")) is not None), runs[0] if runs else None)
    for r in runs:
        if r is not keep:
            r.getparent().remove(r)
    for tag in ("w:hyperlink", "w:ins", "w:smartTag", "w:proofErr", "w:bookmarkStart", "w:bookmarkEnd"):
        for el in list(p.iter(qn(tag))):
            if el.getparent() is not None and (keep is None or keep not in el.iter()):
                el.getparent().remove(el)
    if keep is None:
        keep = p.makeelement(qn("w:r"), {})
        p.append(keep)
    for child in list(keep):
        if child.tag != qn("w:rPr"):
            keep.remove(child)
    t = keep.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    keep.append(t)


def _set_cell(tc, *lines: str):
    paragraphs = tc.findall(qn("w:p"))
    for i, line in enumerate(lines):
        if i >= len(paragraphs):
            new = copy.deepcopy(paragraphs[-1])
            paragraphs[-1].addnext(new)
            paragraphs.append(new)
        _set_text(paragraphs[i], line)
    for extra in paragraphs[max(len(lines), 1):]:
        if _text(extra).strip():   # the template's blank spacing lines stay
            tc.remove(extra)


def _set_header(tc, *lines: str):
    """Header text laid out like the template cell: on separate lines only where the template splits it."""
    if sum(1 for p in tc.findall(qn("w:p")) if _text(p).strip()) > 1:
        _set_cell(tc, *lines)
    else:
        _set_cell(tc, " ".join(lines))


def _cells(tr) -> list:
    return tr.findall(qn("w:tc"))


_AFTER_COLOR = {qn(f"w:{tag}") for tag in (
    "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
    "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath")}


def _set_color(tc, color: Optional[str]):
    """Colours every run in the cell; None leaves the text in the default colour."""
    for r in tc.iter(qn("w:r")):
        r_pr = r.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, r_pr)
        existing = r_pr.find(qn("w:color"))
        if existing is not None:
            r_pr.remove(existing)
        if color:
            el = r_pr.makeelement(qn("w:color"), {qn("w:val"): color})
            # schema order: w:color comes before spacing / size / underline etc.
            after = next((c for c in r_pr if c.tag in _AFTER_COLOR), None)
            if after is not None:
                after.addprevious(el)
            else:
                r_pr.append(el)


def _favour_color(diff: float, cost: bool) -> Optional[str]:
    if abs(diff) < 0.5:
        return None
    return FAVOURABLE if (diff < 0 if cost else diff > 0) else UNFAVOURABLE


def _size_like_cell(tc):
    """Gives the text the cell's own font size (its paragraph mark's).

    The (A) Group table's placeholders are 12pt while its cells, and the (B) and (C) Group tables, are 9pt;
    at 12pt a figure such as 26,367 is wider than the column and wraps.
    """
    for p in tc.findall(qn("w:p")):
        mark = p.find(qn("w:pPr") + "/" + qn("w:rPr"))
        if mark is None:
            continue
        for size in (mark.find(qn("w:sz")), mark.find(qn("w:szCs"))):
            if size is None:
                continue
            for run_size in p.findall(f"{qn('w:r')}/{qn('w:rPr')}/{size.tag}"):
                run_size.set(qn("w:val"), size.get(qn("w:val")))


def _set_column_widths(tbl, widths: Sequence[int]):
    """Fixed column widths in twips (grid, every cell, and the table width)."""
    for col, width in zip(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(width))
    for tr in tbl.findall(qn("w:tr")):
        col = 0
        for tc in _cells(tr):
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is None:
                tc_pr = tc.makeelement(qn("w:tcPr"), {})
                tc.insert(0, tc_pr)
            span_el = tc_pr.find(qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = tc_pr.makeelement(qn("w:tcW"), {})
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(sum(widths[col:col + span])))
            tc_w.set(qn("w:type"), "dxa")
            col += span
    tbl_pr = tbl.find(qn("w:tblPr"))
    for tag, attrs in (("w:tblW", {qn("w:w"): str(sum(widths)), qn("w:type"): "dxa"}),
                       ("w:tblLayout", {qn("w:type"): "fixed"})):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
        tbl_pr.append(tbl_pr.makeelement(qn(tag), attrs))


# ---------------------------------------------------------------------------
# P&L Row Mappings


IS_ROWS = {
    "revenue": lambda l: l["revenue"],
    "opex": lambda l: -l["operating_costs"],
    "ebitda": lambda l: l["ebitda"],
    "operatingprofit": lambda l: l["operating_profit"],
    "operatingprofitloss": lambda l: l["operating_profit"],
    "pbt": lambda l: l["profit_before_tax"],
    "pat": lambda l: l["profit_after_tax"],
}


def _growth(cur: float, base: float):
    diff = cur - base
    return diff, (diff / abs(base) if base else None)


def _margin_change(cur: Optional[float], base: Optional[float]) -> Optional[float]:
    return cur - base if cur is not None and base is not None else None


# ---------------------------------------------------------------------------
# Table Filling Helpers


def _fill_income_statement(tbl, comparison: Comparison, spec: SectionSpec):
    rows = tbl.findall(qn("w:tr"))
    head = _cells(rows[0])
    _set_header(head[1], *spec.header_current)
    _set_header(head[2], *spec.header_base)
    cur, base = comparison.current.lines, comparison.base.lines
    for tr in rows[2:]:
        cells = _cells(tr)
        getter = IS_ROWS.get(_norm(_text(cells[0])))
        if getter is None:
            continue
        diff, pct = _growth(getter(cur), getter(base))
        for tc, value in zip(cells[1:], [money(getter(cur)), money(getter(base)), money(diff), percent(pct)]):
            _set_cell(tc, value)
        _set_color(cells[-1], _favour_color(diff, cost=_norm(_text(cells[0])) == "opex"))


def _update_group_table(tbl, comparison: Comparison, spec: SectionSpec, q: Quarter):
    """Updates the 15-column group table with SLT PLC calculated actuals in columns 6..9."""
    rows = tbl.findall(qn("w:tr"))
    head0 = _cells(rows[0])
    if len(head0) >= 11:
        _set_header(head0[1], spec.current_label, "Group")
        _set_header(head0[2], spec.base_label, "Group")
        _set_header(head0[5], spec.current_label, "SLT PLC")
        _set_header(head0[6], spec.base_label, "SLT PLC")
        _set_header(head0[9], spec.current_label, "Mobitel")
        _set_header(head0[10], spec.base_label, "Mobitel")

    cur, base = comparison.current.lines, comparison.base.lines
    for tr in rows[2:]:
        cells = _cells(tr)
        label = _norm(_text(cells[0]))
        getter = IS_ROWS.get(label)
        if getter is None:
            continue
        cur_val, base_val = getter(cur), getter(base)
        diff, pct = _growth(cur_val, base_val)
        for tc, value in zip(cells[6:10], [money(cur_val), money(base_val), money(diff), percent(pct)]):
            _set_cell(tc, value)
            _size_like_cell(tc)
        _set_color(cells[9], _favour_color(diff, cost=(label == "opex")))


def _fill_rows(tbl, header_rows: int, values: List[List[str]], colors: Sequence[Optional[str]] = ()):
    """Replaces the data rows with one copy of a template data row per entry in `values`."""
    rows = tbl.findall(qn("w:tr"))
    protos = rows[header_rows:]
    middle, last = protos[min(1, len(protos) - 1)], protos[-1]
    for tr in protos:
        tbl.remove(tr)
    for i, row_values in enumerate(values):
        tr = copy.deepcopy(last if i == len(values) - 1 else middle)
        for tc, value in zip(_cells(tr), row_values):
            _set_cell(tc, value)
        if i < len(colors):
            _set_color(_cells(tr)[-1], colors[i])
        tbl.append(tr)


def _driver_rows(comparison: Comparison, line: str, cost: bool):
    """Table rows and their % colours for the accounts that moved a line the most."""
    rows, colors = [], []
    for d in comparison.drivers(line, limit=DRIVER_ROWS):
        sign = -1 if cost else 1
        cur, base = sign * d.current, sign * d.base
        diff, pct = _growth(cur, base)
        rows.append([f"{d.account} {d.name.strip()}", money(cur), money(base), money(diff), percent(pct)])
        colors.append(_favour_color(diff, cost))
    return rows, colors


def _update_financial_summary(tbl, ytd: Comparison, end: Period, ly_end: Period):
    """Fills the P&L lines supported by the trial balance, keeping other rows intact."""
    rows = tbl.findall(qn("w:tr"))
    head = _cells(rows[1])
    _set_cell(head[1], _long(end, "'"))
    _set_cell(head[2], _long(ly_end, "'"))
    cur, base = ytd.current, ytd.base
    amounts = {
        "revenue": "revenue", "operatingprofit": "operating_profit",
        "profitbeforetax": "profit_before_tax", "profitaftertax": "profit_after_tax",
    }
    margins = {
        "ebitdamargin": "ebitda", "operatingprofitmargin": "operating_profit",
        "pbtmargin": "profit_before_tax", "patmargin": "profit_after_tax",
    }
    for tr in rows[2:]:
        cells = _cells(tr)
        label = _norm(_text(cells[0]))
        if not label:
            continue
        if label in amounts:
            key = amounts[label]
            _, pct = _growth(cur.lines[key], base.lines[key])
            _set_cell(cells[1], money(cur.lines[key]))
            _set_cell(cells[2], money(base.lines[key]))
            _set_cell(cells[3], percent(pct))
        elif label in margins:
            key = margins[label]
            c, b = cur.margin(key), base.margin(key)
            _set_cell(cells[1], percent(c))
            _set_cell(cells[2], percent(b))
            _set_cell(cells[3], percent(_margin_change(c, b)))


def _update_key_ratios(tbl, ytd: Comparison, end: Period, ly_end: Period):
    """Updates key ratios supported by the trial balance, keeping other rows intact."""
    rows = tbl.findall(qn("w:tr"))
    head = _cells(rows[2])[1:]
    for tc, text in zip(head, [_long(end, "'"), _long(ly_end, "'"), _long(end, "'")]):
        _set_cell(tc, text)
    cur, base = ytd.current, ytd.base
    _, growth = _growth(cur.lines["revenue"], base.lines["revenue"])

    def coverage(pl):
        interest = -pl.lines["interest_expenses"]
        return f"{pl.lines['operating_profit'] / interest:.1f}" if interest > 0 else "N/A"

    ratios = {
        "revenuegrowthyoy": [percent(growth)],
        "operatingprofitmargin": [percent(cur.margin("operating_profit")), percent(base.margin("operating_profit"))],
        "netprofitmargin": [percent(cur.margin("profit_after_tax")), percent(base.margin("profit_after_tax"))],
        "interestcoverageratiotimes": [coverage(cur), coverage(base)],
    }
    for tr in rows[3:]:
        cells = _cells(tr)
        values = ratios.get(_norm(_text(cells[0])))
        if values:
            for tc, value in zip(cells[1:], values):
                _set_cell(tc, value)


def _update_cover(body_elements, q: Quarter):
    """Updates '2nd Quarter 2026' on the cover to this quarter."""
    for el in body_elements:
        for p in el.iter(qn("w:p")):
            if not re.fullmatch(r"\s*\d(st|nd|rd|th)\s+Quarter\s+\d{4}\s*", _text(p)):
                continue
            runs = [r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
            texts = [r.find(qn("w:t")) for r in runs]
            if len(runs) >= 3 and re.fullmatch(r"st|nd|rd|th", texts[1].text or ""):
                texts[0].text, texts[1].text = str(q.quarter), _ordinal(q.quarter)
                texts[2].text = f" Quarter {q.year}"
                texts[2].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                for t in texts[3:]:
                    t.text = ""
            else:
                _set_text(p, f"{q.quarter}{_ordinal(q.quarter)} Quarter {q.year}")


def _update_executive_summary_slt(doc: Document, yoy: Comparison, phrase: str, spec: SectionSpec):
    paras = doc.paragraphs
    exec_idx = -1
    for i, p in enumerate(paras):
        if _norm(p.text) == "executivesummary":
            exec_idx = i
            break
    if exec_idx == -1:
        return
    slt_idx = -1
    for i in range(exec_idx, len(paras)):
        if _norm(paras[i].text) == "sltplc":
            slt_idx = i
            break
    if slt_idx == -1:
        return
    mob_idx = len(paras)
    for i in range(slt_idx + 1, len(paras)):
        if _norm(paras[i].text) == "mobitel":
            mob_idx = i
            break

    target_paras = [paras[k] for k in range(slt_idx + 1, mob_idx) if paras[k].text.strip()]
    if len(target_paras) >= 1:
        _set_text(target_paras[0]._element, commentary.revenue_paragraph(yoy, ENTITY, phrase))
    if len(target_paras) >= 2:
        _set_text(target_paras[1]._element, commentary.opex_paragraph(yoy, ENTITY, phrase))
    if len(target_paras) >= 3:
        _set_text(target_paras[2]._element, " ".join(commentary.profit_paragraphs(yoy, spec.current_label, spec.base_label)))


def _update_section_commentary(doc: Document, letter: str, comparison: Comparison, spec: SectionSpec):
    paras = doc.paragraphs
    target_pattern = {
        "A": "performanceofthecurrentquarteragainsttheprecedingyear",
        "B": "performanceofthecurrentquarteragainsttheprecedingquarter",
        "C": "performanceofthecurrentyearagainsttheprecedingyear",
    }.get(letter)
    if not target_pattern:
        return
    sec_idx = -1
    for i, p in enumerate(paras):
        if target_pattern in _norm(p.text):
            sec_idx = i
            break
    if sec_idx == -1:
        return
    slt_idx = -1
    for i in range(sec_idx, len(paras)):
        if _norm(paras[i].text) == "sltplc":
            slt_idx = i
            break
    if slt_idx == -1:
        return
    end_idx = len(paras)
    for i in range(slt_idx + 1, len(paras)):
        t = _norm(paras[i].text)
        if t in ("mobitel", "liquidityandoperatingcashflow") or ("performanceof" in t and "performanceofthecurrent" in t):
            end_idx = i
            break

    subheads = {}
    for k in range(slt_idx, end_idx):
        t = _norm(paras[k].text)
        if t in ("revenue", "opex", "profits"):
            subheads[t] = k

    if "revenue" in subheads:
        rev_k = subheads["revenue"]
        next_k = subheads.get("opex", end_idx)
        for k in range(rev_k + 1, next_k):
            if paras[k].text.strip():
                _set_text(paras[k]._element, commentary.revenue_paragraph(comparison, ENTITY, spec.period_phrase))
                break
    if "opex" in subheads:
        opx_k = subheads["opex"]
        next_k = subheads.get("profits", end_idx)
        for k in range(opx_k + 1, next_k):
            if paras[k].text.strip():
                _set_text(paras[k]._element, commentary.opex_paragraph(comparison, ENTITY, spec.period_phrase))
                break
    if "profits" in subheads:
        prf_k = subheads["profits"]
        for k in range(prf_k + 1, end_idx):
            if paras[k].text.strip():
                _set_text(paras[k]._element, " ".join(commentary.profit_paragraphs(comparison, spec.current_label, spec.base_label)))
                break


def _update_section_headings(doc: Document, q: Quarter, specs: List[SectionSpec]):
    """Updates the section headings and table captions with current quarter labels."""
    paras = doc.paragraphs
    for p in paras:
        t = _norm(p.text)
        if "performanceofthecurrentquarteragainsttheprecedingyear" in t:
            _set_text(p._element, f"Performance of the current quarter against the preceding year corresponding quarter "
                                  f"({specs[0].current_label} Vs {specs[0].base_label})")
        elif "performanceofthecurrentquarteragainsttheprecedingquarter" in t:
            _set_text(p._element, f"(B) Performance of the current quarter against the preceding quarter "
                                  f"({specs[1].current_label} vs {specs[1].base_label})")
        elif "performanceofthecurrentyearagainsttheprecedingyear" in t:
            _set_text(p._element, f"Performance of the current year against the preceding year "
                                  f"(YTD {_long(q.end)} Vs YTD {_long(q.same_quarter_last_year().end)})")
        elif "performanceofthecurrentquarteractualsagainstthebudget" in t:
            _set_text(p._element, f"(D) Performance of the current quarter Actuals against the Budget "
                                  f"({_q_short(q)} Actuals Vs Budget)")


# ---------------------------------------------------------------------------
# Pagination
#
# Finance's report has no page breaks: sections reach the top of a page by runs of blank lines sized to
# the template's own text. Generated commentary is longer or shorter, so those runs would end mid-page
# (or spill a page). Each run becomes a page break before the section it pushed down.

SPACER_LINES = 3          # a run this long ends a page in the template; shorter runs are spacing
KEEP_TABLE_ROWS = 30      # tables up to this size fit one page and are kept together
_PPR_FIRST = ("pStyle", "keepNext", "keepLines", "pageBreakBefore")   # schema order at the start of w:pPr


def _set_ppr_flag(p, name: str):
    p_pr = p.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = p.makeelement(qn("w:pPr"), {})
        p.insert(0, p_pr)
    if p_pr.find(qn(f"w:{name}")) is not None:
        return
    el = p_pr.makeelement(qn(f"w:{name}"), {})
    earlier = {qn(f"w:{t}") for t in _PPR_FIRST[:_PPR_FIRST.index(name)]}
    previous = [c for c in p_pr if c.tag in earlier]
    if previous:
        previous[-1].addnext(el)
    else:
        p_pr.insert(0, el)


def _is_spacer(el) -> bool:
    if el.tag != qn("w:p") or _text(el).strip():
        return False
    content = ("w:drawing", "w:pict", "w:object", "w:sectPr", "w:br")
    return not any(next(el.iter(qn(tag)), None) is not None for tag in content)


def _first_paragraph(el):
    return el if el.tag == qn("w:p") else next(el.iter(qn("w:p")))


def _keep_template_pagination(doc: Document):
    body = doc.element.body
    children = [el for el in body if el.tag in (qn("w:p"), qn("w:tbl"))]
    start = next((i for i, el in enumerate(children) if _norm(_text(el)) == "executivesummary"), 0)
    run = []
    for el in children[start:]:
        if _is_spacer(el):
            run.append(el)
            continue
        if len(run) >= SPACER_LINES:
            for spacer in run:
                body.remove(spacer)
            _set_ppr_flag(_first_paragraph(el), "pageBreakBefore")
        run = []

        if el.tag == qn("w:tbl"):
            _keep_table_together(el)
        elif _is_heading(el):
            _set_ppr_flag(el, "keepNext")      # stays with what it introduces


def _keep_table_together(tbl):
    rows = tbl.findall(qn("w:tr"))
    if len(rows) > KEEP_TABLE_ROWS:
        return
    for tr in rows[:-1]:
        for p in tr.iter(qn("w:p")):
            _set_ppr_flag(p, "keepNext")


def _is_heading(p) -> bool:
    """A short line of text (heading or table caption) – not a sentence, a chart or a blank line."""
    text = _text(p).strip()
    has_drawing = next(p.iter(qn("w:drawing")), None) is not None
    return bool(text) and len(text) < 100 and not text.endswith(".") and not has_drawing


# ---------------------------------------------------------------------------
# Template & Builder


class _Template:
    """Locates the paragraphs and tables of Finance's report that the generator copies."""

    def __init__(self, doc):
        self.doc = doc
        self.body = doc.element.body
        children = [el for el in self.body if el.tag in (qn("w:p"), qn("w:tbl"))]
        self.children = children

        exec_idx = self._index(lambda el: _norm(_text(el)) == "executivesummary", "the 'Executive Summary' heading")
        self.cover = children[:exec_idx]
        self.title = children[exec_idx]
        entity_idx = self._index(lambda el: _norm(_text(el)) == "sltplc", "the 'SLT PLC' heading", exec_idx)
        self.entity = children[entity_idx]
        self.body_text = self._find(lambda el: el.tag == qn("w:p") and len(_text(el)) > 80,
                                    "a body paragraph", entity_idx)
        stats_idx = self._index(lambda el: _norm(_text(el)).startswith("financialstatistics"),
                                "the 'Financial Statistics' heading")
        self.stats_title = children[stats_idx]
        self.fin_table = self._find(lambda el: el.tag == qn("w:tbl"), "the Financial Summary table", stats_idx)
        self.ratio_table = self._find(lambda el: el.tag == qn("w:tbl") and el is not self.fin_table,
                                      "the Key Ratios table", stats_idx)
        self.section = self._find(lambda el: _norm(_text(el)).startswith("bperformanceof"),
                                  "the '(B) Performance of ...' heading")
        caption_idx = self._index(lambda el: _norm(_text(el)).startswith("incomestatementsltplc"),
                                  "the 'Income Statement – SLT PLC' caption")
        self.caption = children[caption_idx]
        is_idx = self._index(lambda el: el.tag == qn("w:tbl"), "the SLT PLC income statement table", caption_idx)
        self.is_table = children[is_idx]
        self.subhead = self._find(lambda el: _norm(_text(el)) == "revenue", "the 'Revenue' sub-heading", is_idx)

    def _index(self, predicate: Callable, what: str, start: int = 0) -> int:
        for i in range(start, len(self.children)):
            if predicate(self.children[i]):
                return i
        raise ReportTemplateError(f"Report template is missing {what}")

    def _find(self, predicate: Callable, what: str, start: int = 0):
        return self.children[self._index(predicate, what, start)]


class _Builder:
    def __init__(self, template_path: Path):
        self.doc = Document(str(template_path))
        self.t = _Template(self.doc)
        self.sect_pr = self.doc.element.body.find(qn("w:sectPr"))

    def _append(self, el):
        self.sect_pr.addprevious(el)
        return el

    def paragraph(self, proto, text: str):
        p = copy.deepcopy(proto)
        for drawing in list(p.iter(qn("w:drawing"))):
            drawing.getparent().remove(drawing)
        _set_text(p, text)
        return self._append(p)

    def title(self, text):
        return self.paragraph(self.t.title, text)

    def entity(self, text=ENTITY):
        return self.paragraph(self.t.entity, text)

    def body(self, text):
        return self.paragraph(self.t.body_text, text)

    def subhead(self, text):
        return self.paragraph(self.t.subhead, text)

    def section(self, text):
        return self.paragraph(self.t.section, text)

    def caption(self, text):
        return self.paragraph(self.t.caption, text)

    def picture(self, stream, width_cm: float = 16.0):
        p = self.paragraph(self.t.body_text, "")
        for r in list(p.iter(qn("w:r"))):
            r.getparent().remove(r)
        para = Paragraph(p, self.doc._body)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(stream, width=Cm(width_cm))

    def table(self, proto):
        return self._append(copy.deepcopy(proto))

    def page_break(self):
        p = copy.deepcopy(self.t.body_text)
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        run = p.makeelement(qn("w:r"), {})
        p.append(run)
        br = run.makeelement(qn("w:br"), {qn("w:type"): "page"})
        run.append(br)
        return self._append(p)

    def finish(self, path: Path):
        self.doc.save(str(path))


# ---------------------------------------------------------------------------
# Main Report Generator


def generate_report(inputs: ReportInputs, output_dir: Path, template_path: Path = TEMPLATE_PATH) -> Path:
    analysis = inputs.analysis
    q = analysis.quarter
    specs = _section_specs(q)
    b = _Builder(template_path)
    doc = b.doc

    # 1. Update Cover
    _update_cover(list(doc.element.body), q)

    # 2. Update Executive Summary
    yoy = analysis.comparisons["quarter_vs_last_year"]
    _update_executive_summary_slt(doc, yoy, specs[0].period_phrase, specs[0])

    # 3. Update Financial Statistics
    ytd = analysis.comparisons["ytd_vs_last_year"]
    _update_financial_summary(doc.tables[0]._element, ytd, q.end, q.same_quarter_last_year().end)
    _update_key_ratios(doc.tables[1]._element, ytd, q.end, q.same_quarter_last_year().end)

    # 4. Update Section A (Current Quarter vs Same Quarter Last Year)
    _update_group_table(doc.tables[4]._element, yoy, specs[0], q)
    _fill_income_statement(doc.tables[5]._element, yoy, specs[0])
    _update_section_commentary(doc, "A", yoy, specs[0])

    # 5. Update Section B (Current Quarter vs Preceding Quarter)
    q_vs_prev = analysis.comparisons["quarter_vs_previous"]
    _update_group_table(doc.tables[6]._element, q_vs_prev, specs[1], q)
    _fill_income_statement(doc.tables[7]._element, q_vs_prev, specs[1])
    _update_section_commentary(doc, "B", q_vs_prev, specs[1])

    # 6. Update Section C (Current Year vs Preceding Year YTD)
    _update_group_table(doc.tables[8]._element, ytd, specs[2], q)
    _fill_income_statement(doc.tables[9]._element, ytd, specs[2])
    _update_section_commentary(doc, "C", ytd, specs[2])

    # Update date headers in Section C tables (Cashflow, Borrowings, Interest)
    if len(doc.tables) > 10:
        head_cf = _cells(doc.tables[10]._element.findall(qn("w:tr"))[0])
        if len(head_cf) >= 3:
            _set_cell(head_cf[1], f"YTD {_short(q.end)}")
            _set_cell(head_cf[2], f"YTD {_short(q.same_quarter_last_year().end)}")
    if len(doc.tables) > 12:
        head_int = _cells(doc.tables[12]._element.findall(qn("w:tr"))[0])
        if len(head_int) >= 3:
            _set_header(head_int[1], f"YTD {_short(q.end)}", "(LKR Mn)")
            _set_header(head_int[2], f"YTD {_short(q.same_quarter_last_year().end)}", "(LKR Mn)")

    # 7. Update Section D (Actuals vs Budget)
    if len(doc.tables) > 13:
        t13 = doc.tables[13]._element
        head13 = _cells(t13.findall(qn("w:tr"))[0])
        if len(head13) >= 3:
            _set_header(head13[1], _q_short(q), "(Actuals)")
            _set_header(head13[2], _q_short(q), "(Budget)")
        cur_q_lines = yoy.current.lines
        for tr in t13.findall(qn("w:tr"))[2:]:
            cells = _cells(tr)
            getter = IS_ROWS.get(_norm(_text(cells[0])))
            if getter is not None:
                _set_cell(cells[1], money(getter(cur_q_lines)))

    # Update section headings
    _update_section_headings(doc, q, specs)
    _keep_template_pagination(doc)

    # 8. Append Annexure (ii) – Account-level variance drivers
    b.page_break()
    b.title("Annexure (ii) – Account-level variance drivers")
    for spec in specs:
        comparison = analysis.comparisons[spec.comparison_key]
        for line, name, cost in (("revenue", "Revenue", False), ("operating_costs", "Opex", True)):
            rows, colors = _driver_rows(comparison, line, cost)
            if not rows:
                continue
            b.caption(f"{name} – largest movements by GL account, {spec.current_label} Vs {spec.base_label} (LKR Mn)")
            tbl = b.table(b.t.is_table)
            head = _cells(tbl.findall(qn("w:tr"))[0])
            _set_cell(head[0], "GL account")
            _set_cell(head[1], *spec.header_current)
            _set_cell(head[2], *spec.header_base)
            _fill_rows(tbl, 2, rows, colors)
            _set_column_widths(tbl, DRIVER_COLUMNS)
            b.body("")

    b.title("Sections to be completed by Finance")
    b.body("The following parts of the quarterly report are not in the trial balance and are not generated:")
    for item in MANUAL_SECTIONS:
        b.body(f"•  {item}")

    b.title("Data notes")
    notes = list(inputs.data_notes)
    notes.append("Budget figures have not been loaded, so section (D) shows the template's budget "
                 "placeholders until a budget file is provided.")
    for u in analysis.unmapped:
        notes.append(f"Account {u.account} {u.name.strip()} ({u.period}, LKR {u.amount:,.1f} Mn) "
                     "is not in the GL mapping and is excluded from the P&L.")
    notes.append("Figures are in LKR Mn from the detail trial balances, mapped to P&L lines with the GL code "
                 "mapping. Income is shown positive; Opex is shown as a positive cost. Quarter figures are "
                 "quarter-end YTD less the previous quarter-end YTD.")
    for note in notes:
        b.body(f"•  {note}")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"Financial_Analysis_Q{q.quarter}_{q.year}_{uuid.uuid4().hex[:8]}.docx"
    b.finish(path)
    logger.info(f"Financial Analysis Report written: {path.name}")
    return path
