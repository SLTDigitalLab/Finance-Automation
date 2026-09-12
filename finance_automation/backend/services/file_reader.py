import re
from pathlib import Path
from typing import Tuple, Optional, List
import pandas as pd
import openpyxl
from utils.logger import logger
from config import settings


def detect_header_row(ws, max_scan: int = 30) -> int:
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1
    ):
        row_text = [str(c).strip().lower() if c else "" for c in row]
        if any("account" in t for t in row_text) and any(
            "description" in t for t in row_text
        ):
            return row_idx
        if any("account" in t for t in row_text) and any(
            "balance" in t for t in row_text
        ):
            return row_idx
        if any("account" in t for t in row_text) and any(
            "activity" in t for t in row_text
        ):
            return row_idx
    return 1


def detect_tb_month_and_year(tb_path: str) -> Tuple[str, int]:
    filename = Path(tb_path).stem

    month_map = {
        "jan": "January",
        "feb": "February",
        "mar": "March",
        "apr": "April",
        "may": "May",
        "jun": "June",
        "jul": "July",
        "aug": "August",
        "sep": "September",
        "oct": "October",
        "nov": "November",
        "dec": "December",
    }

    filename_lower = filename.lower()
    detected_month = ""
    detected_year = 0

    # ------------------------
    # First try filename
    # ------------------------
    for key, name in month_map.items():
        if key in filename_lower:
            detected_month = name
            break

    year_match = re.search(r"20\d{2}", filename)
    if year_match:
        detected_year = int(year_match.group())

    # If found from filename, return immediately
    if detected_month and detected_year:
        logger.info(f"Detected TB period from filename: {detected_month} {detected_year}")
        return detected_month, detected_year

    ext = Path(tb_path).suffix.lower()

    # ------------------------
    # TXT fallback
    # ------------------------
    if ext == ".txt":
        try:
            with open(tb_path, "r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    text = line.lower()

                    if not detected_month:
                        for key, name in month_map.items():
                            if key in text:
                                detected_month = name
                                break

                    yr = re.search(r"20\d{2}", line)
                    if yr:
                        detected_year = int(yr.group())

                    if detected_month and detected_year:
                        break

        except Exception as e:
            logger.warning(f"Could not detect month/year from TXT file: {e}")

    # ------------------------
    # Excel fallback
    # ------------------------
    else:
        try:
            wb = openpyxl.load_workbook(tb_path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]

            for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
                for cell in row:
                    if cell is None:
                        continue

                    cell_text = str(cell).lower()

                    if not detected_month:
                        for key, name in month_map.items():
                            if key in cell_text:
                                detected_month = name

                    yr = re.search(r"20\d{2}", str(cell))
                    if yr:
                        detected_year = int(yr.group())

                    if detected_month and detected_year:
                        break

                if detected_month and detected_year:
                    break

            wb.close()

        except Exception as e:
            logger.warning(f"Could not detect month/year from Excel file: {e}")

    logger.info(f"Detected TB period: {detected_month} {detected_year}")

    return detected_month, detected_year


def read_current_year_tb(tb_path: str) -> pd.DataFrame:

    if Path(tb_path).suffix.lower() == ".txt":
        return read_txt_trial_balance(tb_path)

    wb = openpyxl.load_workbook(tb_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = detect_header_row(ws)
    logger.info(f"CY TB header detected at row {header_row}")

    headers = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = [str(c).strip() if c else f"col_{i}" for i, c in enumerate(row)]
    logger.info(f"CY TB headers: {headers}")

    data_rows = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if all(c is None for c in row):
            continue
        first_vals = [str(c).strip() if c else "" for c in row[:3]]
        if all(v == "" or v == "None" for v in first_vals):
            total_row = row[3:] if len(row) > 3 else []
            non_empty = [v for v in total_row if v is not None and str(v).strip()]
            if len(non_empty) <= 1:
                continue

        data_rows.append(row)

    wb.close()

    if not data_rows:
        logger.warning("No data rows found in CY TB")
        return pd.DataFrame(columns=headers)

    df = pd.DataFrame(data_rows)

    num_cols = len(df.columns)
    logger.info(f"CY TB has {num_cols} columns, headers: {headers}")

    flex_col_idx = None
    acct_col_idx = None
    desc_col_idx = None
    bb_col_idx = None
    pa_col_idx = None
    eb_col_idx = None

    for i in range(num_cols):
        h = headers[i].strip().lower() if i < len(headers) else ""
        if "description" in h:
            desc_col_idx = i
        elif "beginning" in h or "opening" in h:
            bb_col_idx = i
        elif "period" in h and "activity" in h:
            pa_col_idx = i
        elif "ending" in h or "closing" in h:
            eb_col_idx = i

    for i in range(num_cols):
        if i in [desc_col_idx, bb_col_idx, pa_col_idx, eb_col_idx]:
            continue
        sample = df.iloc[:, i].dropna().astype(str).head(10)
        has_flex = any("." in str(v) and len(str(v).split(".")) >= 7 for v in sample)
        if has_flex:
            flex_col_idx = i
            break

    for i in range(num_cols):
        if i in [flex_col_idx, desc_col_idx, bb_col_idx, pa_col_idx, eb_col_idx]:
            continue
        sample = df.iloc[:, i].dropna().astype(str).head(5)
        has_digits_only = all(
            re.match(r"^\d{4,6}$", str(v).strip()) for v in sample if str(v).strip()
        )
        if has_digits_only:
            acct_col_idx = i
            break

    if flex_col_idx is None and num_cols >= 3:
        for i in range(num_cols):
            if i in [acct_col_idx, desc_col_idx, bb_col_idx, pa_col_idx, eb_col_idx]:
                continue
            sample = df.iloc[:, i].dropna().astype(str).head(5)
            if any("." in str(v) for v in sample):
                flex_col_idx = i
                break

    logger.info(
        f"CY TB column mapping: acct_idx={acct_col_idx}, desc_idx={desc_col_idx}, "
        f"flex_idx={flex_col_idx}, bb_idx={bb_col_idx}, pa_idx={pa_col_idx}, eb_idx={eb_col_idx}"
    )

    records = []
    for idx, row in df.iterrows():
        gl_code = (
            str(row.iloc[acct_col_idx]).strip() if acct_col_idx is not None else ""
        )
        description = (
            str(row.iloc[desc_col_idx]).strip() if desc_col_idx is not None else ""
        )
        flexfield = (
            str(row.iloc[flex_col_idx]).strip() if flex_col_idx is not None else ""
        )

        if gl_code == "nan":
            gl_code = ""
        if description == "nan":
            description = ""
        if flexfield == "nan":
            flexfield = ""

        if not flexfield or "." not in flexfield:
            continue

        bb = _parse_numeric(row.iloc[bb_col_idx]) if bb_col_idx is not None else 0.0
        pa = _parse_numeric(row.iloc[pa_col_idx]) if pa_col_idx is not None else 0.0
        eb = _parse_numeric(row.iloc[eb_col_idx]) if eb_col_idx is not None else 0.0

        segments = flexfield.split(".")
        while len(segments) < 9:
            segments.append("0")

        records.append(
            {
                "gl_code": gl_code,
                "description": description,
                "flexfield": flexfield,
                "cost_center": segments[1] if len(segments) > 1 else "",
                "location": segments[2] if len(segments) > 2 else "",
                "business_line": segments[3] if len(segments) > 3 else "",
                "product": segments[4] if len(segments) > 4 else "",
                "account": segments[5] if len(segments) > 5 else "",
                "technology": segments[6] if len(segments) > 6 else "",
                "intercompany": segments[7] if len(segments) > 7 else "",
                "project": segments[8] if len(segments) > 8 else "",
                "beginning_balance": bb,
                "period_activity": pa,
                "ending_balance": eb,
                "row_index": idx,
            }
        )

    result_df = pd.DataFrame(records)
    logger.info(f"CY TB loaded: {len(result_df)} valid rows")
    print("\n============= EXCEL DEBUG =============")

    print("Rows :", len(result_df))

    print(result_df.head(10))

    print("Period Activity :", result_df["period_activity"].sum())
    print("Beginning Balance :", result_df["beginning_balance"].sum())
    print("Ending Balance :", result_df["ending_balance"].sum())
    
    print(result_df.iloc[0][["beginning_balance",
                         "period_activity",
                         "ending_balance"]])
    return result_df


def read_previous_year_tb(tb_path: str) -> pd.DataFrame:

    if Path(tb_path).suffix.lower() == ".txt":
        return read_txt_trial_balance(tb_path)

    wb = openpyxl.load_workbook(tb_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = detect_header_row(ws)
    logger.info(f"PY TB header detected at row {header_row}")

    headers = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = [str(c).strip() if c else f"col_{i}" for i, c in enumerate(row)]
    logger.info(f"PY TB headers: {headers}")

    data_rows = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if all(c is None for c in row):
            continue
        non_none = [c for c in row if c is not None]
        if len(non_none) < 2:
            continue
        first_vals = [str(c).strip() if c else "" for c in row[:2]]
        if all(v == "" or v == "None" for v in first_vals):
            numeric_vals = row[2:] if len(row) > 2 else []
            non_empty = [v for v in numeric_vals if v is not None]
            if len(non_empty) <= 1:
                continue

        data_rows.append(row)

    wb.close()

    if not data_rows:
        logger.warning("No data rows found in PY TB")
        return pd.DataFrame()

    df = pd.DataFrame(data_rows)

    num_cols = len(df.columns)
    logger.info(f"PY TB has {num_cols} columns, headers: {headers}")

    desc_col_idx = None
    flex_col_idx = None
    bb_col_idx = None
    pa_col_idx = None
    eb_col_idx = None

    for i in range(num_cols):
        h = headers[i].strip().lower() if i < len(headers) else ""
        if "description" in h:
            desc_col_idx = i
        elif "beginning" in h or "opening" in h:
            bb_col_idx = i
        elif "period" in h and "activity" in h:
            pa_col_idx = i
        elif "ending" in h or "closing" in h:
            eb_col_idx = i

    for i in range(num_cols):
        if i in [desc_col_idx, bb_col_idx, pa_col_idx, eb_col_idx]:
            continue
        sample = df.iloc[:, i].dropna().astype(str).head(10)
        if any("." in str(v) and len(str(v).split(".")) >= 7 for v in sample):
            flex_col_idx = i
            break

    if flex_col_idx is None:
        for i in range(num_cols):
            if i in [desc_col_idx, bb_col_idx, pa_col_idx, eb_col_idx]:
                continue
            sample = df.iloc[:, i].dropna().astype(str).head(10)
            if any(re.match(r"^\d{2}\.\d{4}", str(v)) for v in sample):
                flex_col_idx = i
                break

    logger.info(
        f"PY TB column mapping: desc_idx={desc_col_idx}, flex_idx={flex_col_idx}, "
        f"bb_idx={bb_col_idx}, pa_idx={pa_col_idx}, eb_idx={eb_col_idx}"
    )

    records = []
    for idx, row in df.iterrows():
        description = (
            str(row.iloc[desc_col_idx]).strip() if desc_col_idx is not None else ""
        )
        flexfield = (
            str(row.iloc[flex_col_idx]).strip() if flex_col_idx is not None else ""
        )

        if description == "nan":
            description = ""
        if flexfield == "nan":
            flexfield = ""

        if not flexfield or "." not in flexfield:
            continue

        bb = _parse_numeric(row.iloc[bb_col_idx]) if bb_col_idx is not None else 0.0
        pa = _parse_numeric(row.iloc[pa_col_idx]) if pa_col_idx is not None else 0.0
        eb = _parse_numeric(row.iloc[eb_col_idx]) if eb_col_idx is not None else 0.0

        segments = flexfield.split(".")
        while len(segments) < 9:
            segments.append("0")

        gl_from_flex = segments[5] if len(segments) > 5 else ""

        records.append(
            {
                "gl_code": gl_from_flex,
                "description": description,
                "flexfield": flexfield,
                "cost_center": segments[1] if len(segments) > 1 else "",
                "location": segments[2] if len(segments) > 2 else "",
                "business_line": segments[3] if len(segments) > 3 else "",
                "product": segments[4] if len(segments) > 4 else "",
                "account": segments[5] if len(segments) > 5 else "",
                "technology": segments[6] if len(segments) > 6 else "",
                "intercompany": segments[7] if len(segments) > 7 else "",
                "project": segments[8] if len(segments) > 8 else "",
                "beginning_balance": bb,
                "period_activity": pa,
                "ending_balance": eb,
                "row_index": idx,
            }
        )

    result_df = pd.DataFrame(records)
    logger.info(f"PY TB loaded: {len(result_df)} valid rows")
    
    print("\n============= EXCEL DEBUG =============")

    print("Rows :", len(result_df))

    print(result_df.head(10))

    print("Period Activity :", result_df["period_activity"].sum())
    print("Beginning Balance :", result_df["beginning_balance"].sum())
    print("Ending Balance :", result_df["ending_balance"].sum())
    return result_df

def read_txt_trial_balance(tb_path: str) -> pd.DataFrame:
    records = []

    with open(tb_path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:

            # Remove page-break character
            line = line.replace("\x0c", "").strip()

            if not line:
                continue

            # Skip report headers
            if (
                line.startswith("SLT Primary Ledger")
                or line.startswith("Account")
                or line.startswith("-----------")
                or line.startswith("Currency:")
                or line.startswith("Ledger:")
                or line.startswith("Company:")
                or line.startswith("Company Range:")
                or line.startswith("Report Date:")
                or line.startswith("Page:")
                or "Detail Trial Balance" in line
                or "Year to date" in line
            ):
                continue

            parts = re.split(r"\s{2,}", line)

            if len(parts) < 5:
                continue

            try:
                gl_code = parts[0].strip()
                description = parts[1].strip()
                flexfield = parts[2].strip()

                beginning_balance = _parse_numeric(parts[3])
                period_activity = _parse_numeric(parts[4])

                if len(parts) >= 6:
                    ending_balance = _parse_numeric(parts[5])
                else:
                    ending_balance = beginning_balance + period_activity

                if "." not in flexfield:
                    continue

                segments = flexfield.split(".")

                while len(segments) < 9:
                    segments.append("0")

                # Filter revenue codes (400000 to 429999 inclusive)
                code_val = gl_code if gl_code else segments[5]
                try:
                    code_num = int(float(str(code_val).strip()))
                except (ValueError, TypeError):
                    continue

                if not (400000 <= code_num <= 429999):
                    continue

                records.append(
                    {
                        "gl_code": gl_code,
                        "description": description,
                        "flexfield": flexfield,
                        "cost_center": segments[1],
                        "location": segments[2],
                        "business_line": segments[3],
                        "product": segments[4],
                        "account": segments[5],
                        "technology": segments[6],
                        "intercompany": segments[7],
                        "project": segments[8],
                        "beginning_balance": beginning_balance,
                        "period_activity": period_activity,
                        "ending_balance": ending_balance,
                    }
                )

            except Exception:
                continue

    df = pd.DataFrame(records)
    if not df.empty:
        print(df.iloc[0][["beginning_balance",
                      "period_activity",
                      "ending_balance"]])

    logger.info(f"TXT Trial Balance loaded: {len(df)} rows")

    print("\n================ TXT DEBUG ================")
    print(f"Total rows : {len(df)}")

    if not df.empty:
        print("\nFirst 10 rows")
        print(df.head(10))

        print("\nPeriod Activity Total :", df["period_activity"].sum())
        print("Beginning Balance Total :", df["beginning_balance"].sum())
        print("Ending Balance Total :", df["ending_balance"].sum())

    return df

def _parse_numeric(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    text = text.replace(",", "").replace(" ", "")
    text = text.replace("'", "")
    try:
        return float(text)
    except (ValueError, TypeError):
        return 0.0
