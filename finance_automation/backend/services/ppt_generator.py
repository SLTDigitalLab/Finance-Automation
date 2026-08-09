from pathlib import Path
from typing import List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR_TYPE
from models import RevenueResult
from utils.logger import logger
from utils.formatters import format_number_bracket
from config import settings

# ─── Slide dimensions ────────────────────────────────────────────────────────
SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# ─── Colour palette ───────────────────────────────────────────────────────────
BRAND_BLUE    = RGBColor(0x00, 0x1A, 0x9C)  # #001A9C – title heading
ACCENT_LINE   = RGBColor(0x00, 0x70, 0xC0)  # Blue separator under title
DARK_TEXT     = RGBColor(0x0B, 0x30, 0x41)  # Body / label text

HEADER_FILL   = RGBColor(0xFF, 0xFF, 0x99)  # #FFFF99 – pastel yellow header
DATA_ROW_FILL = RGBColor(0xD9, 0xE1, 0xF2)  # #D9E1F2 – light blue zebra rows
SUBTOTAL_FILL = RGBColor(0x1F, 0x38, 0x64)  # #1F3864 – grand-total dark navy
SUBTOTAL_TEXT = RGBColor(0xFF, 0xFF, 0xFF)  # White text on dark grand-total
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)  # Plain white zebra rows

# ─── Typography ───────────────────────────────────────────────────────────────
FONT_NAME        = "Arial"
TITLE_SIZE       = Pt(28)
HEADER_FONT_SIZE = Pt(10)
DATA_FONT_SIZE   = Pt(9)
FOOTER_FONT_SIZE = Pt(9)

# ─── Logo path ────────────────────────────────────────────────────────────────
_SEEK_LOGO      = Path(__file__).resolve().parents[2] / "frontend" / "src" / "Assets" / "sltmobitel-logo-png_seeklogo-433080.png"
_TEMPLATE_LOGO  = settings.TEMPLATE_DIR / "slt-mobitel-logo.png"
LOGO_PATH = _SEEK_LOGO if _SEEK_LOGO.exists() else _TEMPLATE_LOGO


# ═════════════════════════════════════════════════════════════════════════════
#  Public entry point
# ═════════════════════════════════════════════════════════════════════════════

def generate_pptx(
    revenue_data: List[RevenueResult],
    report_month: str,
    report_year: int,
    output_dir: str,
) -> str:
    logger.info(f"Generating PPTX for {report_month} {report_year}")

    prs = Presentation()
    prs.slide_width  = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide_layout = prs.slide_layouts[6]          # blank layout
    slide = prs.slides.add_slide(slide_layout)

    # ── Slide chrome ──────────────────────────────────────────────────────────
    _add_title(slide)
    _add_unit_label(slide)
    _add_footer(slide, page_num=1)

    # ── Data table ────────────────────────────────────────────────────────────
    num_rows = len(revenue_data) + 2             # 2 header rows + data rows
    num_cols = 11

    table_shape = slide.shapes.add_table(
        num_rows,
        num_cols,
        Inches(0.3),
        Inches(1.1),
        Inches(12.7),
        Inches(5.5),
    )
    table = table_shape.table

    # Column widths: wider label col, equal numeric cols
    table.columns[0].width = Inches(2.2)
    for i in range(1, 11):
        table.columns[i].width = Inches(1.05)

    _set_header_rows(table, report_month, report_year)
    _populate_data_rows(table, revenue_data)
    _style_total_rows(table, revenue_data)

    output_filename = f"Revenue_{report_month}_{report_year}.pptx"
    output_path = Path(output_dir) / output_filename
    prs.save(str(output_path))

    logger.info(f"PPTX saved to: {output_path}")
    return str(output_path)


# ═════════════════════════════════════════════════════════════════════════════
#  Slide chrome
# ═════════════════════════════════════════════════════════════════════════════

def _add_title(slide):
    """Bold brand-blue h1 title + blue separator line beneath."""
    txBox = slide.shapes.add_textbox(
        Inches(0.4), Inches(0.10), Inches(10), Inches(0.60)
    )
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Summary \u2013 Revenue (Fixed)"
    run.font.bold      = True
    run.font.size      = TITLE_SIZE
    run.font.color.rgb = BRAND_BLUE
    run.font.name      = FONT_NAME

    # Horizontal accent line just below the title
    line = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(0.4),   Inches(0.80),
        Inches(12.93), Inches(0.80),
    )
    line.line.color.rgb = ACCENT_LINE
    line.line.width = Pt(1.5)


def _add_unit_label(slide):
    """Right-aligned < Rs. Mn > label above the table."""
    txBox = slide.shapes.add_textbox(
        Inches(10.5), Inches(0.83), Inches(2.1), Inches(0.25)
    )
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "< Rs. Mn >"
    run.font.name      = FONT_NAME
    run.font.size      = Pt(9)
    run.font.italic    = True
    run.font.color.rgb = DARK_TEXT


def _add_footer(slide, page_num: int = 1):
    """Page number centred at bottom; SLTMOBITEL logo pinned bottom-right."""
    footer_y = Inches(7.15)

    # Page number
    pg_box = slide.shapes.add_textbox(
        Inches(6.3), footer_y, Inches(0.5), Inches(0.28)
    )
    tf = pg_box.text_frame
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text           = str(page_num)
    run.font.name      = FONT_NAME
    run.font.size      = FOOTER_FONT_SIZE
    run.font.color.rgb = DARK_TEXT

    # SLTMOBITEL logo – bottom-right corner
    if LOGO_PATH.exists():
        logo_w = Inches(1.65)
        logo_h = Inches(0.36)
        logo_x = SLIDE_WIDTH  - logo_w - Inches(0.12)
        logo_y = SLIDE_HEIGHT - logo_h - Inches(0.04)

        slide.shapes.add_picture(str(LOGO_PATH), logo_x, logo_y, logo_w, logo_h)
    else:
        logger.warning(f"Logo not found at {LOGO_PATH}; skipped.")


# ═════════════════════════════════════════════════════════════════════════════
#  Table header rows
# ═════════════════════════════════════════════════════════════════════════════

def _set_header_rows(table, report_month: str, report_year: int):
    """
    Row 0 (top): blank | Month-MMM'YY (×4) | 20YY Month Act | YTD MMM'YY (×4) | 20YY YTD Act
    Row 1 (sub): blank | Act Bud Vari Vari% | (merged)       | Act Bud Vari Vari%| (merged)
    """
    short_month = report_month[:3] if report_month else "Mon"
    py_year     = report_year - 1

    # Col 0 – label column (merge rows 0+1)
    table.cell(0, 0).merge(table.cell(1, 0))
    _set_cell(table, 0, 0, "", fill_color=HEADER_FILL)

    # Month group (cols 1–4)
    table.cell(0, 1).merge(table.cell(0, 4))
    _set_cell(
        table, 0, 1,
        f"Month \u2013 {short_month} \u2018{str(report_year)[-2:]}",
        bold=True, fill_color=HEADER_FILL, font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER, font_size=HEADER_FONT_SIZE,
    )

    # PY Month (col 5, merge rows 0+1)
    table.cell(0, 5).merge(table.cell(1, 5))
    _set_cell(
        table, 0, 5,
        f"{py_year}\nMonth\nAct",
        bold=True, fill_color=HEADER_FILL, font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER, font_size=HEADER_FONT_SIZE,
    )

    # YTD group (cols 6–9)
    table.cell(0, 6).merge(table.cell(0, 9))
    _set_cell(
        table, 0, 6,
        f"YTD {short_month} \u2018{str(report_year)[-2:]}",
        bold=True, fill_color=HEADER_FILL, font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER, font_size=HEADER_FONT_SIZE,
    )

    # PY YTD (col 10, merge rows 0+1)
    table.cell(0, 10).merge(table.cell(1, 10))
    _set_cell(
        table, 0, 10,
        f"{py_year}\nYTD\nAct",
        bold=True, fill_color=HEADER_FILL, font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER, font_size=HEADER_FONT_SIZE,
    )

    # Sub-header row (row 1): Act / Bud / Vari (underlined) / Vari % (underlined)
    for ci, txt, ul in [
        (1, "Act",    False),
        (2, "Bud",    False),
        (3, "Vari",   True),
        (4, "Vari %", True),
        (6, "Act",    False),
        (7, "Bud",    False),
        (8, "Vari",   True),
        (9, "Vari %", True),
    ]:
        _set_cell(
            table, 1, ci, txt,
            bold=True, fill_color=HEADER_FILL, font_color=DARK_TEXT,
            alignment=PP_ALIGN.CENTER, font_size=HEADER_FONT_SIZE,
            underline=ul,
        )

    # Ensure full header area is yellow (fills any merged remnants)
    for r in range(2):
        for c in range(11):
            _set_cell_fill(table.cell(r, c), HEADER_FILL)


# ═════════════════════════════════════════════════════════════════════════════
#  Data rows
# ═════════════════════════════════════════════════════════════════════════════

def _populate_data_rows(table, revenue_data: List[RevenueResult]):
    """
    Zebra-stripe every non-grand-total row (regular AND subtotals):
      even index → light blue (#D9E1F2)
      odd  index → white
    Subtotal rows use the same stripe colour but with bold text.
    Grand total → full dark navy + white bold text.
    """
    stripe_idx = 0   # increments for every non-grand-total row

    for row_idx, data in enumerate(revenue_data):
        excel_row = row_idx + 2
        is_tot    = data.is_subtotal
        is_grand  = data.is_grand_total

        # ── Choose fills & text colour ─────────────────────────────────────
        if is_grand:
            label_fill = SUBTOTAL_FILL
            data_fill  = SUBTOTAL_FILL
            txt_color  = SUBTOTAL_TEXT
        else:
            # Zebra stripe applied to BOTH label col and data cols
            cell_fill  = DATA_ROW_FILL if (stripe_idx % 2 == 0) else WHITE
            label_fill = cell_fill
            data_fill  = cell_fill
            txt_color  = DARK_TEXT
            stripe_idx += 1

        # ── Label column (col 0) ──────────────────────────────────────────
        _set_cell(
            table, excel_row, 0,
            data.category,
            bold=is_tot, font_color=txt_color,
            alignment=PP_ALIGN.LEFT,
            fill_color=label_fill,
        )

        # ── Numeric columns (1–10) ────────────────────────────────────────
        numeric_cells = [
            (1,  format_number_bracket(data.month_actual)),
            (2,  format_number_bracket(data.month_budget)),
            (3,  format_number_bracket(data.month_variance)),
            (4,  _format_var_pct_display(data.month_variance_pct)),
            (5,  format_number_bracket(data.py_month_actual)),
            (6,  format_number_bracket(data.ytd_actual)),
            (7,  format_number_bracket(data.ytd_budget)),
            (8,  format_number_bracket(data.ytd_variance)),
            (9,  _format_var_pct_display(data.ytd_variance_pct)),
            (10, format_number_bracket(data.py_ytd_actual)),
        ]
        for ci, val in numeric_cells:
            _set_cell(
                table, excel_row, ci,
                val,
                bold=is_tot, font_color=txt_color,
                alignment=PP_ALIGN.RIGHT,
                fill_color=data_fill,
            )


def _style_total_rows(table, revenue_data: List[RevenueResult]):
    """Final pass: bold text on subtotals; full dark-navy on grand total."""
    for row_idx, data in enumerate(revenue_data):
        if not data.is_subtotal:
            continue
        excel_row = row_idx + 2
        for c in range(11):
            cell = table.cell(excel_row, c)
            if data.is_grand_total:
                _set_cell_fill(cell, SUBTOTAL_FILL)
                for para in cell.text_frame.paragraphs:
                    para.font.bold = True
                    para.font.color.rgb = SUBTOTAL_TEXT
            else:
                # Keep zebra fill; just enforce bold dark text
                for para in cell.text_frame.paragraphs:
                    para.font.bold = True
                    para.font.color.rgb = DARK_TEXT


# ═════════════════════════════════════════════════════════════════════════════
#  Cell helpers
# ═════════════════════════════════════════════════════════════════════════════

def _set_cell(
    table,
    row: int,
    col: int,
    text,
    bold: bool = False,
    font_color: Optional[RGBColor] = None,
    fill_color: Optional[RGBColor] = None,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
    font_size=None,
    underline: bool = False,
):
    cell = table.cell(row, col)
    cell.text = ""

    if text is not None and str(text) != "":
        lines = str(text).split("\n")
        for line_idx, line_text in enumerate(lines):
            p = cell.text_frame.paragraphs[0] if line_idx == 0 else cell.text_frame.add_paragraph()
            p.alignment = alignment
            run = p.add_run()
            run.text           = line_text
            run.font.name      = FONT_NAME
            run.font.size      = font_size or DATA_FONT_SIZE
            run.font.bold      = bold
            run.font.underline = underline
            if font_color:
                run.font.color.rgb = font_color

    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left   = Emu(45720)
    cell.margin_right  = Emu(45720)
    cell.margin_top    = Emu(0)
    cell.margin_bottom = Emu(0)

    if fill_color:
        _set_cell_fill(cell, fill_color)


def _set_cell_fill(cell, color: RGBColor):
    from pptx.oxml.ns import qn
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(existing)
    solidFill = tcPr.makeelement(qn("a:solidFill"), {})
    srgbClr   = solidFill.makeelement(qn("a:srgbClr"), {"val": f"{color}"})
    solidFill.append(srgbClr)
    tcPr.append(solidFill)


# ═════════════════════════════════════════════════════════════════════════════
#  Formatting helpers
# ═════════════════════════════════════════════════════════════════════════════

def _format_var_pct_display(value: Optional[float]) -> str:
    if value is None:
        return "#DIV/0!"
    if abs(value) >= 1000:
        return f"{value:,.0f}%"
    if value == int(value):
        return f"{int(value)}%"
    return f"{value:.2f}%"
