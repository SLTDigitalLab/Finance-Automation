import re
import zipfile

import pytest
from docx import Document

from services.financial_analysis.mapping_reader import read_gl_mapping
from services.financial_analysis.periods import Period, Quarter
from services.financial_analysis.pl_engine import build_analysis, monthly_trend
from services.financial_analysis.report_generator import ReportInputs, generate_report
from tests.financial_analysis.helpers import balances, write_mapping

MN = 1_000_000


def ytd(revenue, opex, fx=0.0):
    return balances({411101: -revenue * MN, 511101: opex * MN, 526101: 10 * MN,
                     527201: -fx * MN, 621001: 5 * MN})


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("report")
    mapping = read_gl_mapping(write_mapping(tmp / "map.xlsx"))
    loaded = {
        Period(2025, 3): ytd(100, 60),
        Period(2025, 9): ytd(300, 180),
        Period(2025, 12): ytd(420, 250, fx=4),
        Period(2026, 3): ytd(110, 75, fx=-3),     # Q1 2026: revenue up, opex up, FX loss
    }
    q = Quarter(2026, 1)
    analysis = build_analysis(q, loaded.get, mapping)
    path = generate_report(ReportInputs(analysis, monthly_trend(q.end, loaded.get, mapping), ["Test note"]), tmp)
    return path


def _xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf8"), z.namelist()


def test_cover_shows_the_quarter(report):
    xml, _ = _xml(report)
    texts = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    assert "1st Quarter 2026" in texts
    assert "2nd Quarter" not in texts


def test_reference_document_structure_preserved(report):
    doc = Document(str(report))
    assert len(doc.tables) >= 17
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Executive Summary" in text
    assert "SLT Group" in text
    assert "SLT PLC" in text
    assert "Mobitel" in text
    assert "Key Projects Info" in text or "FTTH Port Utilization" in text
    assert "Annexure" in text


def _income_statements(doc):
    return [t for t in doc.tables
            if len(t.columns) == 5 and t.rows[0].cells[-1].text == "Growth" and t.rows[2].cells[0].text == "Revenue"][:3]


def test_income_statement_tables(report):
    doc = Document(str(report))
    tables = _income_statements(doc)
    assert len(tables) == 3                     # (A), (B), (C)
    q_vs_ly = tables[0]
    assert [c.text for c in q_vs_ly.rows[0].cells[1:3]] == ["Q1\n2026", "Q1\n2025"]
    assert [c.text for c in q_vs_ly.rows[2].cells] == ["Revenue", "110", "100", "10", "10.0%"]
    assert [c.text for c in q_vs_ly.rows[3].cells] == ["Opex", "75", "60", "15", "25.0%"]
    q_vs_q4 = tables[1]
    assert [c.text for c in q_vs_q4.rows[2].cells] == ["Revenue", "110", "120", "(10)", "(8.3%)"]


def test_growth_colours_follow_favourability(report):
    doc = Document(str(report))
    table = _income_statements(doc)[0]

    def color(row):
        run = table.rows[row].cells[-1].paragraphs[0].runs[0]
        return str(run.font.color.rgb)
    assert color(2) == "00B050"     # revenue up: favourable
    assert color(3) == "EE0000"     # opex up: unfavourable


def test_commentary_and_notes(report):
    text = "\n".join(p.text for p in Document(str(report)).paragraphs)
    assert "SLT PLC revenue increased by 10.0% (LKR 10 Mn) in Q1’26 compared to Q1’25" in text
    assert "an exchange loss of LKR 3 Mn in Q1’26, compared to an exchange gain of LKR 4 Mn in Q4’25" in text
    assert "Budget figures have not been loaded" in text
    assert "•  Test note" in text


def test_trend_charts_present(report):
    xml, _ = _xml(report)
    assert xml.count("<pic:pic") >= 3              # cover picture + revenue and opex trend charts


def test_table_formatting_applied(report):
    """Verifies _format_and_paginate_report() was called and applied correctly."""
    from docx.oxml.ns import qn as _qn

    doc = Document(str(report))
    body = doc.element.body
    children = [el for el in body if el.tag in (_qn("w:p"), _qn("w:tbl"), _qn("w:sectPr"))]
    MAX_W = 9570  # page width in dxa with standard margins

    for ti, t in enumerate(doc.tables):
        tbl_el = t._element
        tblPr = tbl_el.find(_qn("w:tblPr"))
        rows = tbl_el.findall(_qn("w:tr"))

        # Every table must be centered
        jc = tblPr.find(_qn("w:jc")) if tblPr is not None else None
        assert jc is not None and jc.get(_qn("w:val")) == "center", (
            f"TABLE#{ti} not centered"
        )

        # No table may exceed the page margins
        tblW_el = tblPr.find(_qn("w:tblW")) if tblPr is not None else None
        if tblW_el is not None:
            w = int(tblW_el.get(_qn("w:w"), "0"))
            assert w <= MAX_W, f"TABLE#{ti} too wide: {w} dxa (max {MAX_W})"

        # Row 0 must have tblHeader
        if rows:
            trPr0 = rows[0].find(_qn("w:trPr"))
            assert trPr0 is not None and trPr0.find(_qn("w:tblHeader")) is not None, (
                f"TABLE#{ti} row 0 missing tblHeader"
            )

        # Short tables (≤20 rows): every data row must have cantSplit
        if 1 < len(rows) <= 20:
            for r_i, tr in enumerate(rows[1:], start=1):
                trPr = tr.find(_qn("w:trPr"))
                assert trPr is not None and trPr.find(_qn("w:cantSplit")) is not None, (
                    f"TABLE#{ti} row {r_i} missing cantSplit"
                )

    # Every non-page-break paragraph immediately before a table must have keepNext
    n = len(children)
    for i, el in enumerate(children):
        if el.tag != _qn("w:p"):
            continue
        if i + 1 < n and children[i + 1].tag == _qn("w:tbl"):
            # Page-break paragraphs are intentionally excluded from keepNext
            has_pb = any(
                br.get(_qn("w:type")) == "page"
                for br in el.iter(_qn("w:br"))
            )
            if has_pb:
                continue
            pPr = el.find(_qn("w:pPr"))
            assert pPr is not None and pPr.find(_qn("w:keepNext")) is not None, (
                "Paragraph immediately before a table missing keepNext: "
                + "".join(t.text or "" for t in el.iter(_qn("w:t")))[:60]
            )

