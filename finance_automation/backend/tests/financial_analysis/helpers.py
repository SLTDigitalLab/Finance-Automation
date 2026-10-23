"""Builds small workbooks in the same layout as the real TB and mapping files."""
import openpyxl

from services.financial_analysis.tb_reader import AccountBalance


def write_tb(path, ytd_header, rows, ptd_header=None, preamble=0):
    """rows: (account, name, beginning, ptd, ytd). account may be str or int."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(preamble):
        ws.append(["Sri Lanka Telecom PLC"])
    ws.append(["A/c", "A/c Name", "GL code", "Beginning balance",
               ptd_header or ytd_header.replace("YTD", "PTD"), ytd_header])
    for acc, name, beg, ptd, ytd in rows:
        gl = f"01.0000.000.00.000.{acc}.00.00.000" if acc is not None else None
        ws.append([acc, name, gl, beg, ptd, ytd])
    wb.save(path)
    return path


def write_mapping(path, extra_rows=()):
    """Mirrors the "6.PL-EBITDA" sheet: labels in B, value in C, From/To in E/F, "and" in G."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "6.PL-EBITDA"
    ws["B3"] = "(all amounts in Sri Lanka Rupees million)"
    ws["C4"], ws["E4"] = "Company-YTD ", "Natural ac codes"
    ws["C5"], ws["E5"], ws["F5"], ws["G5"] = "YTD Sep", "From", "To", "and"
    rows = [
        ("Revenue", 411101, 425501, None),
        ("Operating costs", 511101, 525751, None),
        (None, 525753, 525767, None),
        (None, 525769, 525813, 529999),
        ("EBITDA", None, None, None),
        ("Depreciation", 526101, 526196, None),
        ("Amortisation and impairment of intangible assets", 526202, 526212, None),
        ("Operating profit ", None, None, None),
        ("Other income", 431001, 471099, None),
        (None, 481901, 481999, None),
        ("Interest expenses  ", 527101, 527130, None),
        ("Foreign exchange gain /(loss )", 527201, 527315, None),
        ("Interest income", 481101, 481125, None),
        ("Share of Profit from associates comapny", 525752, None, None),
        ("Profit before tax ", None, None, None),
        ("Taxation", 621001, 621003, None),
        ("Profit After Tax", None, None, None),
        ("Defined benefit plan actuarial (losses)/gain", 525768, None, None),
        ("Tax on comprehensive income", 621004, None, None),
        *extra_rows,
    ]
    for i, (label, start, end, single) in enumerate(rows, start=7):
        ws.cell(i, 2, label)
        ws.cell(i, 3, 123.45)  # cached YTD value column, ignored by the reader
        ws.cell(i, 5, start)
        ws.cell(i, 6, end)
        ws.cell(i, 7, single)
    wb.save(path)
    return path


def balances(ytd_by_account):
    """{account: ytd rupees} -> {account: AccountBalance}"""
    return {acc: AccountBalance(acc, f"Account {acc}", 0.0, 0.0, ytd) for acc, ytd in ytd_by_account.items()}
