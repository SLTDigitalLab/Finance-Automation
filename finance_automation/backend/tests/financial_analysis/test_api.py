"""End-to-end API flow: upload mapping and TBs, generate the Word report, download it."""
import io

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth.jwt_handler import get_current_active_user
from app.database.connection import Base, get_db
from app.models.user_and_log import User
from app.routes import financial_analysis_routes as routes
from tests.financial_analysis.helpers import write_mapping, write_tb

MN = 1_000_000


def _tb(tmp_path, header, revenue, opex):
    """A small balanced TB with YTD revenue / opex in Rs Mn."""
    rows = [
        (411101, "Initiation - Gross", 0, -revenue * 0.3 * MN, -revenue * MN),
        (511101, "Staff Cost", 0, opex * 0.3 * MN, opex * MN),
        (526101, "Depreciation", 0, 0, 10 * MN),
        (111101, "Land", 0, 0, (revenue - opex - 10) * MN),
    ]
    return write_tb(tmp_path / f"{header}.xlsx", header, rows)


@pytest.fixture
def client(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    db = Session()
    admin = User(full_name="Test Admin", email="admin@example.com", role="Admin", status="Approved")
    db.add(admin)
    db.commit()

    monkeypatch.setattr(routes, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(routes, "MAPPING_DIR", tmp_path / "uploads" / "mapping")
    monkeypatch.setattr(routes, "OUTPUT_DIR", tmp_path / "outputs")

    app = FastAPI()
    app.include_router(routes.financial_analysis_router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_active_user] = lambda: admin
    yield TestClient(app)
    db.close()


def _upload(client, url, paths):
    files = [("files" if url.endswith("trial-balances") else "file",
              (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
    return client.post(url, files=files)


def test_full_flow(client, tmp_path):
    api = "/api/financial-analysis"

    resp = client.post(f"{api}/generate", json={"year": 2026, "quarter": 2})
    assert resp.status_code == 422
    assert "mapping" in resp.json()["detail"]

    resp = _upload(client, f"{api}/mapping", [write_mapping(tmp_path / "map.xlsx")])
    assert resp.status_code == 200, resp.text
    assert resp.json()["mapping"]["filename"] == "map.xlsx"

    bad = tmp_path / "notes.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Hello"])
    wb.save(bad)
    tbs = [_tb(tmp_path, "YTD Mar 25", 100, 60), _tb(tmp_path, "YTD Jun 25", 220, 130),
           _tb(tmp_path, "YTD March 26", 150, 90), bad]
    resp = _upload(client, f"{api}/trial-balances", tbs)
    results = resp.json()["results"]
    assert [r["status"] for r in results] == ["loaded", "loaded", "loaded", "error"]
    assert results[2]["period"] == "Mar 2026"

    status = client.get(f"{api}/status").json()
    assert [p["label"] for p in status["periods"]] == ["Mar 2026", "Jun 2025", "Mar 2025"]

    resp = client.post(f"{api}/generate", json={"year": 2026, "quarter": 2})
    assert resp.status_code == 422
    assert "Jun 2026" in resp.json()["detail"]

    resp = _upload(client, f"{api}/trial-balances", [_tb(tmp_path, "YTD June 26", 330, 190)])
    assert resp.json()["results"][0]["status"] == "loaded"
    q2 = next(q for q in client.get(f"{api}/status").json()["quarters"] if q["label"] == "Q2 2026")
    assert q2["missing"] == []

    resp = client.post(f"{api}/generate", json={"year": 2026, "quarter": 2})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    vs_ly = body["comparisons"][0]
    revenue = vs_ly["rows"][0]
    assert vs_ly["title"] == "Q2 2026 vs Q2 2025"
    assert (revenue["current"], revenue["base"], revenue["variance_pct"]) == (180.0, 120.0, 50.0)
    opex = vs_ly["rows"][1]
    assert (opex["label"], opex["current"], opex["base"], opex["favourable"]) == ("Opex", 100.0, 70.0, False)

    resp = client.get(f"{api}/download/{body['filename']}")
    assert resp.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
    assert "Executive Summary" in text
    assert "SLT PLC revenue increased by 50.0% (LKR 60 Mn) in Q2’26 compared to Q2’25" in text

    assert client.get(f"{api}/download/..%2Fsecret.docx").status_code in (400, 404)

    assert client.delete(f"{api}/trial-balances/2026/6").status_code == 200
    assert client.delete(f"{api}/trial-balances/2026/6").status_code == 404
